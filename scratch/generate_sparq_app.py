import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor

slug = "sparq-dotnet_principal_architect-260831"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for .NET Principal Architect)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">.NET Principal Architect &amp; Cloud Modernization Lead</div>
    <div class="contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>.NET Principal Architect with 15+ years leading enterprise application modernization, cloud enablement, and distributed systems architecture. Proven track record assessing legacy .NET Framework monoliths and establishing repeatable patterns for migration to Microsoft Azure and modern .NET 8. Deep expertise in monolithic refactoring, SOAP-to-REST transitions, Windows Services containerization (Docker, AKS), SQL Server performance tuning, and AI-assisted development evangelism.</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">Application Modernization:</span> .NET 8, .NET Core, .NET Framework / ASP.NET, Monolithic Refactoring, Domain-Driven Design (DDD), Clean Architecture, CQRS, Microservices.</p>
  <p class="competency-row"><span class="competency-label">Microsoft Azure &amp; Cloud:</span> Azure App Services, AKS, Azure Functions, Azure SQL / Managed Instance, Networking, IAM, RBAC, Cloud Security Guardrails, GCP.</p>
  <p class="competency-row"><span class="competency-label">Service &amp; Messaging Evolution:</span> SOAP to REST Migration, Windows Services to Containerized Workloads (Docker), MSMQ to Azure Service Bus / Event-Driven Queues, WCF.</p>
  <p class="competency-row"><span class="competency-label">Data &amp; Infrastructure as Code:</span> SQL Server (Schema Design, Query Optimization, Stored Procedures, Indexing), Terraform, Azure DevOps, CI/CD, Git.</p>
  <p class="competency-row"><span class="competency-label">Quality, Leadership &amp; AI:</span> Technical Governance, Code Reviews, TDD, AI-Assisted Engineering (MLOps, GenAI Certified 2026), Multi-Team Mentorship.</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer | Aven Hospitality</span> <span>Jan 2026 - Present</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Remote/Hybrid)</div>
    <ul>
      <li>Lead the modernization roadmap for core hospitality platforms, upgrading legacy .NET Framework backends to modern .NET and cloud services.</li>
      <li>Architected decoupled event-driven messaging pipelines using cloud queues (PubSub / Service Bus) to handle high-volume booking transactions.</li>
      <li>Containerized backend services and automated CI/CD deployment pipelines via Jenkins, PowerShell scripts, and Docker.</li>
      <li>Established architecture guardrails, automated code quality benchmarks, and vulnerability remediation with SonarQube and Veracode.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led an engineering pod of 8 Senior .NET engineers delivering mission-critical microservices and transactional business logic on Microsoft Azure.</li>
      <li>Architected cloud-native services using Azure Kubernetes Service (AKS), Azure Functions, and App Services, cutting transaction latency.</li>
      <li>Defined repeatable migration standards, peer review guidelines, and automated release verification pipelines in Azure DevOps.</li>
      <li>Guided engineers through database performance tuning, SQL Server stored procedure optimization, and resiliency testing.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; .NET Developer | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Architected enterprise cloud solutions and custom .NET extensions for Fortune 500 accounts in the United States and United Kingdom.</li>
      <li>Engineered developer productivity tooling in C# integrated with enterprise AI coding assistants and automation platforms.</li>
      <li>Advised client engineering teams on cloud governance, API management, and scalable integration patterns.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Consultant | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led a team of 4 engineers refactoring core enterprise monoliths for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server.</li>
      <li>Designed enterprise portal solutions and Business Intelligence platforms for major corporate and regulated clients, including Banco de Chile.</li>
      <li>Implemented continuous integration and automated build verification on TeamCity servers.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect &amp; .NET Developer | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Engineered high-availability, self-healing Windows Services, cross-platform APIs, and recursive SQL stored procedures for Envac AB worldwide operations.</li>
      <li>Architected Azure backend REST services and payment gateway integrations for Collector Bank mobile payment platform.</li>
      <li>Led core architectural delivery for Volvo Cars Corporation global enterprise intranet platform on Microsoft technologies.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed multithreaded Windows Services, yield optimization algorithms, and high-throughput transaction backends in C# and ASP.NET.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>Microsoft Certifications:</span> <span>Power Platform Developer Associate (PL-400), App Maker (PL-100), MCTS C# &amp; SharePoint</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; MLOps:</span> <span>MLOps Essentials (2026), Generative AI Pair Programming (2026), Azure DevOps CI/CD</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('.NET Principal Architect & Cloud Modernization Lead')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    '.NET Principal Architect with 15+ years leading enterprise application modernization, cloud enablement, and distributed systems '
    'architecture. Proven track record assessing legacy .NET Framework monoliths and establishing repeatable patterns for migration '
    'to Microsoft Azure and modern .NET 8. Deep expertise in monolithic refactoring, SOAP-to-REST transitions, Windows Services '
    'containerization (Docker, AKS), SQL Server performance tuning, and AI-assisted development evangelism.'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('Application Modernization', '.NET 8, .NET Core, .NET Framework / ASP.NET, Monolithic Refactoring, Domain-Driven Design (DDD), Clean Architecture, CQRS, Microservices.')
add_comp('Microsoft Azure & Cloud', 'Azure App Services, AKS, Azure Functions, Azure SQL / Managed Instance, Networking, IAM, RBAC, Cloud Security Guardrails, GCP.')
add_comp('Service & Messaging Evolution', 'SOAP to REST Migration, Windows Services to Containerized Workloads (Docker), MSMQ to Azure Service Bus / Event-Driven Queues, WCF.')
add_comp('Data & Infrastructure as Code', 'SQL Server (Schema Design, Query Optimization, Stored Procedures, Indexing), Terraform, Azure DevOps, CI/CD, Git.')
add_comp('Quality, Leadership & AI', 'Technical Governance, Code Reviews, TDD, AI-Assisted Engineering (MLOps, GenAI Certified 2026), Multi-Team Mentorship.')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer', 'Aven Hospitality', 'Jan 2026 - Present', 'Montevideo, Uruguay (Remote/Hybrid)',
    [
        'Lead the modernization roadmap for core hospitality platforms, upgrading legacy .NET Framework backends to modern .NET and cloud services.',
        'Architected decoupled event-driven messaging pipelines using cloud queues (PubSub / Service Bus) to handle high-volume booking transactions.',
        'Containerized backend services and automated CI/CD deployment pipelines via Jenkins, PowerShell scripts, and Docker.',
        'Established architecture guardrails, automated code quality benchmarks, and vulnerability remediation with SonarQube and Veracode.'
    ]
)

add_role_cv(
    'Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Led an engineering pod of 8 Senior .NET engineers delivering mission-critical microservices and transactional business logic on Microsoft Azure.',
        'Architected cloud-native services using Azure Kubernetes Service (AKS), Azure Functions, and App Services, cutting transaction latency.',
        'Defined repeatable migration standards, peer review guidelines, and automated release verification pipelines in Azure DevOps.',
        'Guided engineers through database performance tuning, SQL Server stored procedure optimization, and resiliency testing.'
    ]
)

add_role_cv(
    'Power Platform Architect & .NET Developer', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Architected enterprise cloud solutions and custom .NET extensions for Fortune 500 accounts in the United States and United Kingdom.',
        'Engineered developer productivity tooling in C# integrated with enterprise AI coding assistants and automation platforms.',
        'Advised client engineering teams on cloud governance, API management, and scalable integration patterns.'
    ]
)

add_role_cv(
    'Software Architect & Consultant', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led a team of 4 engineers refactoring core enterprise monoliths for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server.',
        'Designed enterprise portal solutions and Business Intelligence platforms for major corporate and regulated clients, including Banco de Chile.',
        'Implemented continuous integration and automated build verification on TeamCity servers.'
    ]
)

add_role_cv(
    'Technical Solution Architect & .NET Developer', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Engineered high-availability, self-healing Windows Services, cross-platform APIs, and recursive SQL stored procedures for Envac AB worldwide operations.',
        'Architected Azure backend REST services and payment gateway integrations for Collector Bank mobile payment platform.',
        'Led core architectural delivery for Volvo Cars Corporation global enterprise intranet platform on Microsoft technologies.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Developed multithreaded Windows Services, yield optimization algorithms, and high-throughput transaction backends in C# and ASP.NET.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('Microsoft Certifications: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('Power Platform Developer Associate (PL-400), App Maker (PL-100), MCTS C# & SharePoint')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('AI & MLOps: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('MLOps Essentials (2026), Generative AI Pair Programming (2026), Azure DevOps CI/CD')

pc3 = doc_cv.add_paragraph()
pc3.paragraph_format.space_after = Pt(2)
rc3_l = pc3.add_run('International Patent: ')
rc3_l.bold = True
rc3_l.font.color.rgb = PRIMARY_COLOR
pc3.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")

# -------------------------------------------------------------
# 3. Cover Letter HTML (English, addressed to Sparq Hiring Team)
# -------------------------------------------------------------
cl_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - Cover Letter</title>
<style>
  @page { size: A4; margin: 2.2cm 2.4cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 11pt; color: #1a1a1a; background: #ffffff; line-height: 1.45; }

  .sender { border-bottom: 2pt solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 14pt; }
  .sender-name { font-size: 15pt; font-weight: bold; color: #1f3a5f; }
  .sender-contact { font-size: 9.5pt; color: #444; margin-top: 2pt; }

  .date { margin-bottom: 14pt; font-size: 10pt; color: #555; }
  .recipient { margin-bottom: 14pt; font-size: 10.5pt; }
  p { margin: 0 0 9pt 0; }
  ul { margin: 0 0 9pt 0; padding-left: 16pt; break-inside: avoid; }
  li { margin-bottom: 3pt; }
  .signature { margin-top: 14pt; }
</style>
</head>
<body>
  <div class="sender">
    <div class="sender-name">Pedro Milburn</div>
    <div class="sender-contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>
  <div class="date">August 31, 2026</div>
  <div class="recipient">Hiring Team / Cloud Enablement<br>Sparq Holdings, Inc. (recruiting@teamsparq.com)</div>

  <p>Dear Sparq Hiring Team,</p>

  <p>I am applying for the <strong>.NET Principal Architect</strong> role on Sparq's Cloud Enablement Team. With 15+ years architecting enterprise applications, assessing legacy codebases, and leading complex cloud modernizations, I have dedicated my career to the exact mission Sparq champions: helping organizations solve the right problems and transition confidently to modern, cloud-native architectures.</p>

  <p>Throughout my consulting and engineering leadership career at Globant, Sabre Corporation, and Aven Hospitality, I have defined repeatable modernization patterns that turn monolithic systems into scalable, secure Azure platforms:</p>

  <ul>
    <li><strong>Monolithic Refactoring &amp; Azure Enablement:</strong> Led the migration of legacy .NET Framework monoliths to modern .NET and Azure microservices (App Services, AKS, Azure Functions), establishing robust guardrails around networking, RBAC, and secure cloud operations.</li>
    <li><strong>Service &amp; Database Evolution:</strong> Guided transitions from legacy SOAP/WCF services and multithreaded Windows Services to lightweight RESTful APIs and containerized Docker workloads; optimized mission-critical SQL Server databases and implemented event-driven messaging (Azure Service Bus).</li>
    <li><strong>Engineering Craftsmanship &amp; AI Acceleration:</strong> Mentored multi-team pods in Clean Architecture, TDD, and Terraform IaC, actively adopting and evangelizing AI-assisted engineering practices (certified in MLOps and Generative AI Pair Programming).</li>
  </ul>

  <p>Based in Uruguay and working seamlessly across North American time zones with full bilingual proficiency in English and Spanish, I bring a collaborative, ego-free approach to technical leadership and client partnerships.</p>

  <p>I look forward to discussing how my cloud modernization experience can accelerate outcomes for Sparq and your clients.</p>

  <div class="signature">Sincerely,<br><strong>Pedro Milburn</strong></div>
</body>
</html>
"""

cl_html_path = app_dir / f"CL-{slug}.html"
cl_html_path.write_text(cl_html.strip(), encoding="utf-8")
print(f"Written CL HTML: {cl_html_path}")

# -------------------------------------------------------------
# 4. Cover Letter DOCX (English)
# -------------------------------------------------------------
cl_docx_path = app_dir / f"CL-{slug}.docx"
doc_cl = docx.Document()

for section in doc_cl.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style_cl = doc_cl.styles['Normal']
style_cl.font.name = 'Arial'
style_cl.font.size = Pt(10.5)
style_cl.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

p_cl_name = doc_cl.add_paragraph()
p_cl_name.paragraph_format.space_after = Pt(2)
r_cl_name = p_cl_name.add_run('Pedro Milburn')
r_cl_name.bold = True
r_cl_name.font.size = Pt(16)
r_cl_name.font.color.rgb = PRIMARY_COLOR

p_cl_contact = doc_cl.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_cl_contact.runs[0].font.size = Pt(9.5)
p_cl_contact.runs[0].font.color.rgb = GRAY_COLOR
p_cl_contact.paragraph_format.space_after = Pt(14)

p_cl_date = doc_cl.add_paragraph('August 31, 2026')
p_cl_date.paragraph_format.space_after = Pt(12)

p_cl_rec = doc_cl.add_paragraph('Hiring Team / Cloud Enablement\nSparq Holdings, Inc. (recruiting@teamsparq.com)')
p_cl_rec.paragraph_format.space_after = Pt(14)

p_cl_sal = doc_cl.add_paragraph('Dear Sparq Hiring Team,')
p_cl_sal.paragraph_format.space_after = Pt(10)

p_cl_open = doc_cl.add_paragraph(
    'I am applying for the .NET Principal Architect role on Sparq\'s Cloud Enablement Team. With 15+ years architecting '
    'enterprise applications, assessing legacy codebases, and leading complex cloud modernizations, I have dedicated my '
    'career to the exact mission Sparq champions: helping organizations solve the right problems and transition '
    'confidently to modern, cloud-native architectures.'
)
p_cl_open.paragraph_format.space_after = Pt(10)
p_cl_open.paragraph_format.line_spacing = 1.15

p_cl_body = doc_cl.add_paragraph(
    'Throughout my consulting and engineering leadership career at Globant, Sabre Corporation, and Aven Hospitality, '
    'I have defined repeatable modernization patterns that turn monolithic systems into scalable, secure Azure platforms:'
)
p_cl_body.paragraph_format.space_after = Pt(8)
p_cl_body.paragraph_format.line_spacing = 1.15

cl_b1 = doc_cl.add_paragraph(style='List Bullet')
r1 = cl_b1.add_run('Monolithic Refactoring & Azure Enablement: ')
r1.bold = True
cl_b1.add_run(
    'Led the migration of legacy .NET Framework monoliths to modern .NET and Azure microservices (App Services, AKS, Azure Functions), '
    'establishing robust guardrails around networking, RBAC, and secure cloud operations.'
)
cl_b1.paragraph_format.space_after = Pt(4)
cl_b1.paragraph_format.line_spacing = 1.15

cl_b2 = doc_cl.add_paragraph(style='List Bullet')
r2 = cl_b2.add_run('Service & Database Evolution: ')
r2.bold = True
cl_b2.add_run(
    'Guided transitions from legacy SOAP/WCF services and multithreaded Windows Services to lightweight RESTful APIs and containerized Docker '
    'workloads; optimized mission-critical SQL Server databases and implemented event-driven messaging (Azure Service Bus).'
)
cl_b2.paragraph_format.space_after = Pt(4)
cl_b2.paragraph_format.line_spacing = 1.15

cl_b3 = doc_cl.add_paragraph(style='List Bullet')
r3 = cl_b3.add_run('Engineering Craftsmanship & AI Acceleration: ')
r3.bold = True
cl_b3.add_run(
    'Mentored multi-team pods in Clean Architecture, TDD, and Terraform IaC, actively adopting and evangelizing AI-assisted engineering '
    'practices (certified in MLOps and Generative AI Pair Programming).'
)
cl_b3.paragraph_format.space_after = Pt(10)
cl_b3.paragraph_format.line_spacing = 1.15

p_cl_close1 = doc_cl.add_paragraph(
    'Based in Uruguay and working seamlessly across North American time zones with full bilingual proficiency in English and Spanish, '
    'I bring a collaborative, ego-free approach to technical leadership and client partnerships.'
)
p_cl_close1.paragraph_format.space_after = Pt(10)
p_cl_close1.paragraph_format.line_spacing = 1.15

p_cl_close2 = doc_cl.add_paragraph('I look forward to discussing how my cloud modernization experience can accelerate outcomes for Sparq and your clients.')
p_cl_close2.paragraph_format.space_after = Pt(16)

p_cl_sign = doc_cl.add_paragraph('Sincerely,\n')
r_cl_sig = p_cl_sign.add_run('Pedro Milburn')
r_cl_sig.bold = True

doc_cl.save(str(cl_docx_path))
print(f"Written CL DOCX: {cl_docx_path}")
