import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor

slug = "missionhires-supervisor_desarrollo-260826"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# 1. HTML Cover Letter
cl_html = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - Carta de Presentación</title>
<style>
  @page { size: A4; margin: 2.2cm 2.4cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 11pt; color: #1a1a1a; background: #ffffff; line-height: 1.45; }

  .sender { border-bottom: 2pt solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 14pt; }
  .sender-name { font-size: 15pt; font-weight: bold; color: #1f3a5f; }
  .sender-contact { font-size: 9.5pt; color: #444; margin-top: 2pt; }

  .date { margin-bottom: 14pt; font-size: 10pt; color: #555; }
  .recipient { margin-bottom: 14pt; font-size: 10.5pt; }
  p { margin: 0 0 9pt 0; }
  ul { margin: 0 0 9pt 0; padding-left: 16pt; break-inside: avoid; }
  li { margin-bottom: 3pt; }
  .signature { margin-top: 14pt; }
</style>
</head>
<body>
  <div class="sender">
    <div class="sender-name">Pedro Milburn</div>
    <div class="sender-contact">Montevideo, Uruguay &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401</div>
  </div>
  <div class="date">26 de agosto de 2026</div>
  <div class="recipient">Mariana Escolar / Equipo de Selección<br>MissionHires</div>

  <p>Estimada Mariana y equipo de selección de MissionHires,</p>

  <p>Me postulo a la posición de <strong>Supervisor/a de Desarrollo para la empresa nacional del sector financiero</strong>. Cuento con más de 15 años de trayectoria en ingeniería de software, arquitectura de soluciones y liderazgo de equipos de desarrollo, especializándome en la conducción de proyectos tecnológicos complejos, gestión de proveedores y gobierno técnico en el ecosistema .NET y SQL Server.</p>

  <p>En mis roles recientes como Lead Software Engineer en Aven Hospitality y Development Team Lead en Sabre Corporation, he liderado y supervisado equipos multidisciplinarios de 8 profesionales (desarrolladores senior y especialistas de QA), planificando bajo marcos ágiles (Scrum/Kanban) y asegurando el estricto cumplimiento de plazos, alcance y estándares de calidad:</p>

  <ul>
    <li><strong>Supervisión de Equipos y Gestión de Entregas:</strong> Experiencia en asignación de tareas, seguimiento de desempeño, estimación de hitos, mitigación proactiva de riesgos y reporte continuo a niveles gerenciales y de dirección.</li>
    <li><strong>Experiencia en Sector Financiero y Entornos Regulados:</strong> Desarrollo e implementación de sistemas bancarios de digitalización y trazabilidad para Banco Itaú en 20 sucursales, pasarelas y servicios de pagos en Azure para Collector Bank, y software de gobierno y cumplimiento normativo ISO en MVD Technologies.</li>
    <li><strong>Excelencia Técnica y Calidad:</strong> Sólido dominio de arquitecturas .NET, bases de datos SQL Server, gestión en Azure DevOps/Jira, control de versiones en Git y aseguramiento de código mediante pruebas automatizadas (TDD) y análisis de seguridad con SonarQube y Veracode.</li>
  </ul>

  <p>Residiendo en Montevideo y con amplia experiencia en articulación entre áreas de negocio, equipos técnicos y proveedores estratégicos, aporto autonomía, criterio de toma de decisiones y compromiso con los objetivos del sector financiero nacional.</p>

  <p>Agradezco su tiempo y consideración, y quedo a su entera disposición para coordinar una entrevista.</p>

  <div class="signature">Atentamente,<br><strong>Pedro Milburn</strong></div>
</body>
</html>
"""

html_path = app_dir / f"CL-{slug}.html"
html_path.write_text(cl_html.strip(), encoding="utf-8")
print(f"Written HTML Cover Letter to {html_path}")

# 2. DOCX Cover Letter
docx_path = app_dir / f"CL-{slug}.docx"
doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

p_name = doc.add_paragraph()
p_name.paragraph_format.space_after = Pt(2)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(16)
r_name.font.color.rgb = RGBColor(0x1f, 0x3a, 0x5f)

p_contact = doc.add_paragraph('Montevideo, Uruguay · pierre.milburn@hotmail.com · +598 92 806 401')
p_contact.runs[0].font.size = Pt(9.5)
p_contact.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
p_contact.paragraph_format.space_after = Pt(14)

p_date = doc.add_paragraph('26 de agosto de 2026')
p_date.paragraph_format.space_after = Pt(12)

p_rec = doc.add_paragraph('Mariana Escolar / Equipo de Selección\nMissionHires')
p_rec.paragraph_format.space_after = Pt(14)

p_sal = doc.add_paragraph('Estimada Mariana y equipo de selección de MissionHires,')
p_sal.paragraph_format.space_after = Pt(10)

p_open = doc.add_paragraph(
    'Me postulo a la posición de Supervisor/a de Desarrollo para la empresa nacional del sector financiero. '
    'Cuento con más de 15 años de trayectoria en ingeniería de software, arquitectura de soluciones y liderazgo '
    'de equipos de desarrollo, especializándome en la conducción de proyectos tecnológicos complejos, gestión de '
    'proveedores y gobierno técnico en el ecosistema .NET y SQL Server.'
)
p_open.paragraph_format.space_after = Pt(10)
p_open.paragraph_format.line_spacing = 1.15

p_body = doc.add_paragraph(
    'En mis roles recientes como Lead Software Engineer en Aven Hospitality y Development Team Lead en Sabre Corporation, '
    'he liderado y supervisado equipos multidisciplinarios de 8 profesionales (desarrolladores senior y especialistas de QA), '
    'planificando bajo marcos ágiles (Scrum/Kanban) y asegurando el estricto cumplimiento de plazos, alcance y estándares de calidad:'
)
p_body.paragraph_format.space_after = Pt(8)
p_body.paragraph_format.line_spacing = 1.15

b1 = doc.add_paragraph(style='List Bullet')
r1_bold = b1.add_run('Supervisión de Equipos y Gestión de Entregas: ')
r1_bold.bold = True
b1.add_run(
    'Experiencia en asignación de tareas, seguimiento de desempeño, estimación de hitos, mitigación proactiva '
    'de riesgos y reporte continuo a niveles gerenciales y de dirección.'
)
b1.paragraph_format.space_after = Pt(4)
b1.paragraph_format.line_spacing = 1.15

b2 = doc.add_paragraph(style='List Bullet')
r2_bold = b2.add_run('Experiencia en Sector Financiero y Entornos Regulados: ')
r2_bold.bold = True
b2.add_run(
    'Desarrollo e implementación de sistemas bancarios de digitalización y trazabilidad para Banco Itaú en 20 sucursales, '
    'pasarelas y servicios de pagos en Azure para Collector Bank, y software de gobierno y cumplimiento normativo ISO en MVD Technologies.'
)
b2.paragraph_format.space_after = Pt(4)
b2.paragraph_format.line_spacing = 1.15

b3 = doc.add_paragraph(style='List Bullet')
r3_bold = b3.add_run('Excelencia Técnica y Calidad: ')
r3_bold.bold = True
b3.add_run(
    'Sólido dominio de arquitecturas .NET, bases de datos SQL Server, gestión en Azure DevOps/Jira, control de versiones '
    'en Git y aseguramiento de código mediante pruebas automatizadas (TDD) y análisis de seguridad con SonarQube y Veracode.'
)
b3.paragraph_format.space_after = Pt(10)
b3.paragraph_format.line_spacing = 1.15

p_close1 = doc.add_paragraph(
    'Residiendo en Montevideo y con amplia experiencia en articulación entre áreas de negocio, equipos técnicos '
    'y proveedores estratégicos, aporto autonomía, criterio de toma de decisiones y compromiso con los objetivos '
    'del sector financiero nacional.'
)
p_close1.paragraph_format.space_after = Pt(10)
p_close1.paragraph_format.line_spacing = 1.15

p_close2 = doc.add_paragraph('Agradezco su tiempo y consideración, y quedo a su entera disposición para coordinar una entrevista.')
p_close2.paragraph_format.space_after = Pt(16)

p_sign = doc.add_paragraph('Atentamente,\n')
r_sig = p_sign.add_run('Pedro Milburn')
r_sig.bold = True

doc.save(str(docx_path))
print(f'Written DOCX Cover Letter to {docx_path}')
