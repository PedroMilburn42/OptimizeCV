import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
import win32com.client

slug = "rurika-full_stack_engineer_ai-260903"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for Rurika Full Stack / Applied AI)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">Full Stack &amp; Applied AI Engineer | Cloud Solutions Architect (C#, Python, Azure)</div>
    <div class="contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>Full Stack Engineer and Cloud Solutions Architect with 15+ years of software engineering experience and deep specialization in Applied AI, Large Language Model (LLM) integration, and cloud-native architectures on Microsoft Azure and GCP. Proven track record taking AI and data experimentation proofs of concept (POCs/MVPs) into secure, production-grade enterprise deployments. Expert in C#, Python, RESTful APIs, MLOps, Responsible AI, and modern AI-assisted engineering workflows (Claude Code, GitHub Copilot).</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">Applied AI &amp; LLM Orchestration:</span> LLM Application Integration, Context Engineering, Prompt Architecture, Gemini, Copilot, MLOps Essentials (2026), Foundations of Responsible AI (2026), AI Pair Programming.</p>
  <p class="competency-row"><span class="competency-label">Cloud Architecture &amp; Azure:</span> Microsoft Azure (App Services, AKS, Azure Functions, Azure Cognitive &amp; ML Services, Logic Apps), GCP (Cloud Run, VMs, PubSub), Terraform, Cloud Security &amp; Threat Modeling.</p>
  <p class="competency-row"><span class="competency-label">Full Stack &amp; Backend Engineering:</span> C#, modern .NET (.NET 8, .NET Core), Python (NLP, ML, scripts), RESTful APIs, SOA, Microservices, JavaScript/TypeScript, Angular, HTML5.</p>
  <p class="competency-row"><span class="competency-label">Data Platforms &amp; Databases:</span> SQL Server (T-SQL, query tuning, stored procedures), PostgreSQL (ETL/ELT), Microsoft Fabric (SQL Databases 2026), MySQL, Oracle, Entity Framework, Dapper.</p>
  <p class="competency-row"><span class="competency-label">DevOps, Security &amp; Quality:</span> CI/CD Automation (Azure DevOps, Jenkins), Docker, Security Peer Reviews (Veracode, SonarQube), TDD, Automated Testing, Open Source Delivery.</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer &amp; AI Architect | Aven Hospitality</span> <span>Jan 2026 - Jul 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Remote/Hybrid)</div>
    <ul>
      <li>Led a technical team of 8 engineers and QA specialists designing scalable cloud components and investigating AI-powered orchestration workflows.</li>
      <li>Architected context pipelines and integrations for LLM-assisted developer tooling using Gemini, Copilot, and CodeNova.</li>
      <li>Migrated legacy services to Google Cloud Platform (GCP Cloud Run, PubSub, VMs), designing event-driven pipelines for high-throughput booking transactions.</li>
      <li>Conducted security code reviews and threat modeling with Veracode and SonarQube, automating CI/CD release pipelines in Jenkins.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Manager &amp; Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Managed an engineering pod of 8 Senior .NET developers delivering core transaction microservices on Microsoft Azure.</li>
      <li>Architected distributed cloud solutions on Azure (AKS, Azure Functions, App Services), resolving critical bottlenecks in transactional latency.</li>
      <li>Collaborated across cross-functional teams (Product, Data, Operations) to define architectural roadmaps and establish repeatable Azure deployment patterns.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; AI Tooling Developer | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Developed AI-powered developer productivity plugins in C# and Visual Studio Code for Globant's internal Augmented Coding initiative.</li>
      <li>Architected and deployed enterprise cloud solutions, RESTful API integrations, and Power Platform automations for Fortune 500 clients in the US and UK.</li>
      <li>Standardized Azure DevOps CI/CD pipelines and established security best practices across distributed engineering pods.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Program Manager | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led 4 delivery teams architecting Business Intelligence solutions, SharePoint portals, and core system refactoring for UPM Uruguay and Banco de Chile.</li>
      <li>Implemented continuous integration, automated unit testing, and decoupled 3-layer architectures with Entity Framework and SQL Server.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect &amp; Cloud Developer | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li><strong>Collector Bank (2014):</strong> Architected Azure Mobile Services backend REST API and AngularJS frontend for mobile invoice payment validation.</li>
      <li><strong>Envac AB (2013-2015):</strong> Developed high-availability Windows Services, cross-platform APIs, and recursive SQL stored procedures for worldwide operations.</li>
      <li><strong>Volvo Cars Corporation (2011-2013):</strong> Delivered core architecture and development for the global worldwide enterprise platform on Microsoft cloud services.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed yield optimization algorithms, multithreaded backend services, and high-throughput data processing systems in C# and ASP.NET.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; MLOps Certifications:</span> <span>Foundations of Responsible AI (2026), MLOps Essentials (2026), AI-Powered Pair Programming (2026), Advanced NLP with Python (2026)</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Microsoft &amp; Cloud Credentials:</span> <span>Microsoft Certified PL-400, PL-100, MCTS C#, Bases de datos SQL en Fabric (2026), Developing CI/CD on Azure (2023)</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248 · registered in US, Europe, Sweden)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX (English)
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('Full Stack & Applied AI Engineer | Cloud Solutions Architect (C#, Python, Azure)')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    'Full Stack Engineer and Cloud Solutions Architect with 15+ years of software engineering experience and deep specialization '
    'in Applied AI, Large Language Model (LLM) integration, and cloud-native architectures on Microsoft Azure and GCP. Proven track '
    'record taking AI and data experimentation proofs of concept (POCs/MVPs) into secure, production-grade enterprise deployments. '
    'Expert in C#, Python, RESTful APIs, MLOps, Responsible AI, and modern AI-assisted engineering workflows (Claude Code, GitHub Copilot).'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('Applied AI & LLM Orchestration', 'LLM Application Integration, Context Engineering, Prompt Architecture, Gemini, Copilot, MLOps Essentials (2026), Foundations of Responsible AI (2026), AI Pair Programming.')
add_comp('Cloud Architecture & Azure', 'Microsoft Azure (App Services, AKS, Azure Functions, Azure Cognitive & ML Services, Logic Apps), GCP (Cloud Run, VMs, PubSub), Terraform, Cloud Security & Threat Modeling.')
add_comp('Full Stack & Backend Engineering', 'C#, modern .NET (.NET 8, .NET Core), Python (NLP, ML, scripts), RESTful APIs, SOA, Microservices, JavaScript/TypeScript, Angular, HTML5.')
add_comp('Data Platforms & Databases', 'SQL Server (T-SQL, query tuning, stored procedures), PostgreSQL (ETL/ELT), Microsoft Fabric (SQL Databases 2026), MySQL, Oracle, Entity Framework, Dapper.')
add_comp('DevOps, Security & Quality', 'CI/CD Automation (Azure DevOps, Jenkins), Docker, Security Peer Reviews (Veracode, SonarQube), TDD, Automated Testing, Open Source Delivery.')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer & AI Architect', 'Aven Hospitality', 'Jan 2026 - Jul 2026', 'Montevideo, Uruguay (Remote/Hybrid)',
    [
        'Led a technical team of 8 engineers and QA specialists designing scalable cloud components and investigating AI-powered orchestration workflows.',
        'Architected context pipelines and integrations for LLM-assisted developer tooling using Gemini, Copilot, and CodeNova.',
        'Migrated legacy services to Google Cloud Platform (GCP Cloud Run, PubSub, VMs), designing event-driven pipelines for high-throughput booking transactions.',
        'Conducted security code reviews and threat modeling with Veracode and SonarQube, automating CI/CD release pipelines in Jenkins.'
    ]
)

add_role_cv(
    'Manager & Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Managed an engineering pod of 8 Senior .NET developers delivering core transaction microservices on Microsoft Azure.',
        'Architected distributed cloud solutions on Azure (AKS, Azure Functions, App Services), resolving critical bottlenecks in transactional latency.',
        'Collaborated across cross-functional teams (Product, Data, Operations) to define architectural roadmaps and establish repeatable Azure deployment patterns.'
    ]
)

add_role_cv(
    'Power Platform Architect & AI Tooling Developer', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Developed AI-powered developer productivity plugins in C# and Visual Studio Code for Globant\'s internal Augmented Coding initiative.',
        'Architected and deployed enterprise cloud solutions, RESTful API integrations, and Power Platform automations for Fortune 500 clients in the US and UK.',
        'Standardized Azure DevOps CI/CD pipelines and established security best practices across distributed engineering pods.'
    ]
)

add_role_cv(
    'Software Architect & Program Manager', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led 4 delivery teams architecting Business Intelligence solutions, SharePoint portals, and core system refactoring for UPM Uruguay and Banco de Chile.',
        'Implemented continuous integration, automated unit testing, and decoupled 3-layer architectures with Entity Framework and SQL Server.'
    ]
)

add_role_cv(
    'Technical Solution Architect & Cloud Developer', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Collector Bank (2014): Architected Azure Mobile Services backend REST API and AngularJS frontend for mobile invoice payment validation.',
        'Envac AB (2013-2015): Developed high-availability Windows Services, cross-platform APIs, and recursive SQL stored procedures for worldwide operations.',
        'Volvo Cars Corporation (2011-2013): Delivered core architecture and development for the global worldwide enterprise platform on Microsoft cloud services.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Developed yield optimization algorithms, multithreaded backend services, and high-throughput data processing systems in C# and ASP.NET.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('AI & MLOps Certifications: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('Foundations of Responsible AI (2026), MLOps Essentials (2026), AI-Powered Pair Programming (2026), Advanced NLP with Python (2026)')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('Microsoft & Cloud Credentials: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('Microsoft Certified PL-400, PL-100, MCTS C#, Bases de datos SQL en Fabric (2026), Developing CI/CD on Azure (2023)')

pc3 = doc_cv.add_paragraph()
pc3.paragraph_format.space_after = Pt(2)
rc3_l = pc3.add_run('International Patent: ')
rc3_l.bold = True
rc3_l.font.color.rgb = PRIMARY_COLOR
pc3.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248 · registered in US, Europe, Sweden)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")

# -------------------------------------------------------------
# 3. Cover Letter HTML (English, addressed to Rurika)
# -------------------------------------------------------------
cl_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - Cover Letter</title>
<style>
  @page { size: A4; margin: 2.2cm 2.4cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 11pt; color: #1a1a1a; background: #ffffff; line-height: 1.45; }

  .sender { border-bottom: 2pt solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 14pt; }
  .sender-name { font-size: 15pt; font-weight: bold; color: #1f3a5f; }
  .sender-contact { font-size: 9.5pt; color: #444; margin-top: 2pt; }

  .date { margin-bottom: 14pt; font-size: 10pt; color: #555; }
  .recipient { margin-bottom: 14pt; font-size: 10.5pt; }
  p { margin: 0 0 9pt 0; }
  ul { margin: 0 0 9pt 0; padding-left: 16pt; break-inside: avoid; }
  li { margin-bottom: 3pt; }
  .signature { margin-top: 14pt; }
</style>
</head>
<body>
  <div class="sender">
    <div class="sender-name">Pedro Milburn</div>
    <div class="sender-contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>
  <div class="date">September 3, 2026</div>
  <div class="recipient">Technical Hiring Team<br>Rurika Consultancy</div>

  <p>Dear Rurika Hiring Team,</p>

  <p>I am writing to express my strong interest in the <strong>Full Stack Engineer (Applied AI &amp; Cloud Platforms)</strong> position at Rurika Consultancy. With 15+ years of software engineering and cloud architecture experience, combined with hands-on expertise building Applied AI solutions, LLM integrations, and robust cloud platforms, I specialize in bridging software craftsmanship with data science and business objectives.</p>

  <p>Rurika's mission to build foundational cloud solutions supporting Applied AI experimentation while maintaining rigorous security and responsible data practices directly reflects my recent work and core capabilities:</p>

  <ul>
    <li><strong>Applied AI &amp; LLM Solution Delivery:</strong> Experience designing context pipelines and orchestration for LLM applications (Gemini, Copilot, CodeNova), certified in Foundations of Responsible AI, MLOps Essentials, and AI-Powered Pair Programming (2026), with a focus on solution quality evaluation and data privacy.</li>
    <li><strong>Cloud Architecture &amp; Data Foundation:</strong> Extensive experience designing scalable cloud architectures on Microsoft Azure (AKS, App Services, Functions, Cognitive Services) and GCP (Cloud Run, PubSub), supported by deep database engineering in SQL Server, PostgreSQL, and Microsoft Fabric.</li>
    <li><strong>POC-to-Production Velocity &amp; Security:</strong> Track record creating rapid proofs of concept and minimum viable products (MVPs) and successfully advancing them to production, conducting threat modeling, security code reviews (Veracode/SonarQube), and automating CI/CD pipelines.</li>
  </ul>

  <p>Working fully remote from Uruguay with complete bilingual fluency in English and Spanish, I bring collaborative thought leadership, clear async technical communication, and a passion for unblocking cross-functional teams.</p>

  <p>I welcome the opportunity to discuss how my technical leadership and Applied AI experience can support Rurika's client initiatives.</p>

  <div class="signature">Sincerely,<br><strong>Pedro Milburn</strong></div>
</body>
</html>
"""

cl_html_path = app_dir / f"CL-{slug}.html"
cl_html_path.write_text(cl_html.strip(), encoding="utf-8")
print(f"Written CL HTML: {cl_html_path}")

# -------------------------------------------------------------
# 4. Cover Letter DOCX (English)
# -------------------------------------------------------------
cl_docx_path = app_dir / f"CL-{slug}.docx"
doc_cl = docx.Document()

for section in doc_cl.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style_cl = doc_cl.styles['Normal']
style_cl.font.name = 'Arial'
style_cl.font.size = Pt(10.5)
style_cl.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

p_cl_name = doc_cl.add_paragraph()
p_cl_name.paragraph_format.space_after = Pt(2)
r_cl_name = p_cl_name.add_run('Pedro Milburn')
r_cl_name.bold = True
r_cl_name.font.size = Pt(16)
r_cl_name.font.color.rgb = PRIMARY_COLOR

p_cl_contact = doc_cl.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_cl_contact.runs[0].font.size = Pt(9.5)
p_cl_contact.runs[0].font.color.rgb = GRAY_COLOR
p_cl_contact.paragraph_format.space_after = Pt(14)

p_cl_date = doc_cl.add_paragraph('September 3, 2026')
p_cl_date.paragraph_format.space_after = Pt(12)

p_cl_rec = doc_cl.add_paragraph('Technical Hiring Team\nRurika Consultancy')
p_cl_rec.paragraph_format.space_after = Pt(14)

p_cl_sal = doc_cl.add_paragraph('Dear Rurika Hiring Team,')
p_cl_sal.paragraph_format.space_after = Pt(10)

p_cl_open = doc_cl.add_paragraph(
    'I am writing to express my strong interest in the Full Stack Engineer (Applied AI & Cloud Platforms) position at Rurika '
    'Consultancy. With 15+ years of software engineering and cloud architecture experience, combined with hands-on expertise building '
    'Applied AI solutions, LLM integrations, and robust cloud platforms, I specialize in bridging software craftsmanship with data '
    'science and business objectives.'
)
p_cl_open.paragraph_format.space_after = Pt(10)
p_cl_open.paragraph_format.line_spacing = 1.15

p_cl_body = doc_cl.add_paragraph(
    'Rurika\'s mission to build foundational cloud solutions supporting Applied AI experimentation while maintaining rigorous security and responsible data practices directly reflects my recent work and core capabilities:'
)
p_cl_body.paragraph_format.space_after = Pt(8)
p_cl_body.paragraph_format.line_spacing = 1.15

cl_b1 = doc_cl.add_paragraph(style='List Bullet')
r1 = cl_b1.add_run('Applied AI & LLM Solution Delivery: ')
r1.bold = True
cl_b1.add_run(
    'Experience designing context pipelines and orchestration for LLM applications (Gemini, Copilot, CodeNova), certified in Foundations of Responsible AI, MLOps Essentials, and AI-Powered Pair Programming (2026), with a focus on solution quality evaluation and data privacy.'
)
cl_b1.paragraph_format.space_after = Pt(4)
cl_b1.paragraph_format.line_spacing = 1.15

cl_b2 = doc_cl.add_paragraph(style='List Bullet')
r2 = cl_b2.add_run('Cloud Architecture & Data Foundation: ')
r2.bold = True
cl_b2.add_run(
    'Extensive experience designing scalable cloud architectures on Microsoft Azure (AKS, App Services, Functions, Cognitive Services) and GCP (Cloud Run, PubSub), supported by deep database engineering in SQL Server, PostgreSQL, and Microsoft Fabric.'
)
cl_b2.paragraph_format.space_after = Pt(4)
cl_b2.paragraph_format.line_spacing = 1.15

cl_b3 = doc_cl.add_paragraph(style='List Bullet')
r3 = cl_b3.add_run('POC-to-Production Velocity & Security: ')
r3.bold = True
cl_b3.add_run(
    'Track record creating rapid proofs of concept and minimum viable products (MVPs) and successfully advancing them to production, conducting threat modeling, security code reviews (Veracode/SonarQube), and automating CI/CD pipelines.'
)
cl_b3.paragraph_format.space_after = Pt(10)
cl_b3.paragraph_format.line_spacing = 1.15

p_cl_close1 = doc_cl.add_paragraph(
    'Working fully remote from Uruguay with complete bilingual fluency in English and Spanish, I bring collaborative thought leadership, '
    'clear async technical communication, and a passion for unblocking cross-functional teams.'
)
p_cl_close1.paragraph_format.space_after = Pt(10)
p_cl_close1.paragraph_format.line_spacing = 1.15

p_cl_close2 = doc_cl.add_paragraph('I welcome the opportunity to discuss how my technical leadership and Applied AI experience can support Rurika\'s client initiatives.')
p_cl_close2.paragraph_format.space_after = Pt(16)

p_cl_sign = doc_cl.add_paragraph('Sincerely,\n')
r_cl_sig = p_cl_sign.add_run('Pedro Milburn')
r_cl_sig.bold = True

doc_cl.save(str(cl_docx_path))
print(f"Written CL DOCX: {cl_docx_path}")

# -------------------------------------------------------------
# 5. Convert DOCX to PDF via Word COM
# -------------------------------------------------------------
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    # CV PDF
    cv_pdf_path = app_dir / f"CV-{slug}.pdf"
    doc1 = word.Documents.Open(str(cv_docx_path.resolve()))
    doc1.SaveAs(str(cv_pdf_path.resolve()), FileFormat=17)
    doc1.Close()
    print(f"Written CV PDF: {cv_pdf_path}")

    # CL PDF
    cl_pdf_path = app_dir / f"CL-{slug}.pdf"
    doc2 = word.Documents.Open(str(cl_docx_path.resolve()))
    doc2.SaveAs(str(cl_pdf_path.resolve()), FileFormat=17)
    doc2.Close()
    print(f"Written CL PDF: {cl_pdf_path}")
finally:
    word.Quit()

print("All Rurika application documents generated successfully!")
