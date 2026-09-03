import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
import win32com.client

slug = "descartes-dotnet_software_developer-260903"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for Descartes .NET / Data Integrations)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">Senior .NET &amp; Data Integrations Developer (C#, SQL Server, Azure)</div>
    <div class="contact">Montevideo, Uruguay (Hybrid) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>Senior .NET Software Developer and Solutions Architect with 15+ years of experience specializing in backend data integrations, high-volume transactional processing, and relational database engineering. Deep expertise in C#, modern .NET, Windows Worker Services, advanced SQL Server (T-SQL, stored procedures, ETL pipelines, triggers, performance optimization), REST APIs, and B2B data transfer protocols (XML, JSON, CSV). Active advocate for AI-first development, automated testing, and CI/CD within Microsoft Azure environments.</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">.NET &amp; Backend Engineering:</span> C#, .NET 8, .NET Core, .NET Framework, Entity Framework Core, LINQ, ADO.NET, Windows Worker Services, Dependency Injection, REST APIs, Microservices.</p>
  <p class="competency-row"><span class="competency-label">Database &amp; SQL Server Expertise:</span> SQL Server (Complex Stored Procedures, Triggers, Functions, T-SQL Tuning, Index Optimization, Execution Plans), ETL/ELT Pipelines, Data Modeling, PostgreSQL, MySQL.</p>
  <p class="competency-row"><span class="competency-label">Data Integrations &amp; Protocols:</span> High-Volume Transaction Processing, Data Ingestion &amp; Validation, FTP / SFTP, XML, JSON, CSV, B2B Communication Layers, Encryption &amp; Security.</p>
  <p class="competency-row"><span class="competency-label">Cloud, DevOps &amp; Quality:</span> Microsoft Azure (App Services, Functions, AKS, Azure SQL), Git, CI/CD (Azure DevOps, Jenkins), PowerShell, Unit Testing (TDD), Veracode, SonarQube.</p>
  <p class="competency-row"><span class="competency-label">AI-First Engineering:</span> AI-Assisted Development (Claude Code, GitHub Copilot), MLOps Essentials (2026), AI-Powered Pair Programming (2026), Process Automation.</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer | Aven Hospitality</span> <span>Jan 2026 - Jul 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Hybrid)</div>
    <ul>
      <li>Led the development and optimization of high-throughput transactional booking backend services in C# and modern .NET.</li>
      <li>Architected decoupled data ingestion pipelines and asynchronous message queues (PubSub) for real-time transaction processing.</li>
      <li>Optimized relational database queries, automated CI/CD deployment pipelines in Jenkins, and maintained zero-downtime releases.</li>
      <li>Enforced rigorous code reviews, automated unit testing, and vulnerability remediation with SonarQube and Veracode.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Manager &amp; Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Managed an engineering pod of 8 Senior .NET developers delivering core transaction logic and data processing microservices on Azure.</li>
      <li>Architected cloud-native services (Azure Functions, App Services, AKS), resolving critical latency bottlenecks in high-volume queries.</li>
      <li>Led database performance tuning, SQL Server stored procedure optimization, and data integrity validation across global travel systems.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; .NET Developer | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Architected and delivered custom .NET cloud extensions, RESTful API integrations, and B2B data pipelines for Fortune 500 accounts in the US and UK.</li>
      <li>Developed developer productivity tooling and AI automation plugins in C# and TypeScript.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Program Manager | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led core architectural refactoring for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server 2014.</li>
      <li>Designed Business Intelligence data pipelines and reporting integrations for Banco de Chile and GNL Quintero.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect &amp; .NET Developer | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li><strong>Envac AB (2013-2015):</strong> Built self-healing, high-availability Windows Services, cross-platform APIs, and recursive SQL stored procedures for worldwide operations.</li>
      <li><strong>Collector Bank (2014):</strong> Developed Azure backend REST API for mobile invoice payment validation and financial data processing.</li>
      <li><strong>Volvo Cars Corporation (2011-2013):</strong> Delivered core architecture and enterprise data integrations on Microsoft platforms.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed multithreaded Windows Services, data yield optimization algorithms, and high-volume data ingestion pipelines in C# and ASP.NET.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>Microsoft &amp; Database Credentials:</span> <span>MCTS: Programming in C# (70-483), Bases de datos SQL en Fabric (2026), ETL/ELT in PostgreSQL (2023), PL-400</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; DevOps:</span> <span>Generative AI Pair Programming (2026), MLOps Essentials (2026), Developing CI/CD on Azure (2023)</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX (English)
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('Senior .NET & Data Integrations Developer (C#, SQL Server, Azure)')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay (Hybrid) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    'Senior .NET Software Developer and Solutions Architect with 15+ years of experience specializing in backend data integrations, '
    'high-volume transactional processing, and relational database engineering. Deep expertise in C#, modern .NET, Windows Worker Services, '
    'advanced SQL Server (T-SQL, stored procedures, ETL pipelines, triggers, performance optimization), REST APIs, and B2B data transfer '
    'protocols (XML, JSON, CSV). Active advocate for AI-first development, automated testing, and CI/CD within Microsoft Azure environments.'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('.NET & Backend Engineering', 'C#, .NET 8, .NET Core, .NET Framework, Entity Framework Core, LINQ, ADO.NET, Windows Worker Services, Dependency Injection, REST APIs, Microservices.')
add_comp('Database & SQL Server Expertise', 'SQL Server (Complex Stored Procedures, Triggers, Functions, T-SQL Tuning, Index Optimization, Execution Plans), ETL/ELT Pipelines, Data Modeling, PostgreSQL, MySQL.')
add_comp('Data Integrations & Protocols', 'High-Volume Transaction Processing, Data Ingestion & Validation, FTP / SFTP, XML, JSON, CSV, B2B Communication Layers, Encryption & Security.')
add_comp('Cloud, DevOps & Quality', 'Microsoft Azure (App Services, Functions, AKS, Azure SQL), Git, CI/CD (Azure DevOps, Jenkins), PowerShell, Unit Testing (TDD), Veracode, SonarQube.')
add_comp('AI-First Engineering', 'AI-Assisted Development (Claude Code, GitHub Copilot), MLOps Essentials (2026), AI-Powered Pair Programming (2026), Process Automation.')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer', 'Aven Hospitality', 'Jan 2026 - Jul 2026', 'Montevideo, Uruguay (Hybrid)',
    [
        'Led the development and optimization of high-throughput transactional booking backend services in C# and modern .NET.',
        'Architected decoupled data ingestion pipelines and asynchronous message queues (PubSub) for real-time transaction processing.',
        'Optimized relational database queries, automated CI/CD deployment pipelines in Jenkins, and maintained zero-downtime releases.',
        'Enforced rigorous code reviews, automated unit testing, and vulnerability remediation with SonarQube and Veracode.'
    ]
)

add_role_cv(
    'Manager & Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Managed an engineering pod of 8 Senior .NET developers delivering core transaction logic and data processing microservices on Azure.',
        'Architected cloud-native services (Azure Functions, App Services, AKS), resolving critical latency bottlenecks in high-volume queries.',
        'Led database performance tuning, SQL Server stored procedure optimization, and data integrity validation across global travel systems.'
    ]
)

add_role_cv(
    'Power Platform Architect & .NET Developer', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Architected and delivered custom .NET cloud extensions, RESTful API integrations, and B2B data pipelines for Fortune 500 accounts in the US and UK.',
        'Developed developer productivity tooling and AI automation plugins in C# and TypeScript.'
    ]
)

add_role_cv(
    'Software Architect & Program Manager', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led core architectural refactoring for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server 2014.',
        'Designed Business Intelligence data pipelines and reporting integrations for Banco de Chile and GNL Quintero.'
    ]
)

add_role_cv(
    'Technical Solution Architect & .NET Developer', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Envac AB (2013-2015): Built self-healing, high-availability Windows Services, cross-platform APIs, and recursive SQL stored procedures for worldwide operations.',
        'Collector Bank (2014): Developed Azure backend REST API for mobile invoice payment validation and financial data processing.',
        'Volvo Cars Corporation (2011-2013): Delivered core architecture and enterprise data integrations on Microsoft platforms.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Developed multithreaded Windows Services, data yield optimization algorithms, and high-volume data ingestion pipelines in C# and ASP.NET.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('Microsoft & Database Credentials: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('MCTS: Programming in C# (70-483), Bases de datos SQL en Fabric (2026), ETL/ELT in PostgreSQL (2023), PL-400')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('AI & DevOps: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('Generative AI Pair Programming (2026), MLOps Essentials (2026), Developing CI/CD on Azure (2023)')

pc3 = doc_cv.add_paragraph()
pc3.paragraph_format.space_after = Pt(2)
rc3_l = pc3.add_run('International Patent: ')
rc3_l.bold = True
rc3_l.font.color.rgb = PRIMARY_COLOR
pc3.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")

# -------------------------------------------------------------
# 3. Cover Letter HTML (English, addressed to Descartes)
# -------------------------------------------------------------
cl_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - Cover Letter</title>
<style>
  @page { size: A4; margin: 2.2cm 2.4cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 11pt; color: #1a1a1a; background: #ffffff; line-height: 1.45; }

  .sender { border-bottom: 2pt solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 14pt; }
  .sender-name { font-size: 15pt; font-weight: bold; color: #1f3a5f; }
  .sender-contact { font-size: 9.5pt; color: #444; margin-top: 2pt; }

  .date { margin-bottom: 14pt; font-size: 10pt; color: #555; }
  .recipient { margin-bottom: 14pt; font-size: 10.5pt; }
  p { margin: 0 0 9pt 0; }
  ul { margin: 0 0 9pt 0; padding-left: 16pt; break-inside: avoid; }
  li { margin-bottom: 3pt; }
  .signature { margin-top: 14pt; }
</style>
</head>
<body>
  <div class="sender">
    <div class="sender-name">Pedro Milburn</div>
    <div class="sender-contact">Montevideo, Uruguay (Hybrid) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>
  <div class="date">September 3, 2026</div>
  <div class="recipient">Hiring Team / Data Integrations<br>Descartes Systems Group (Montevideo)</div>

  <p>Dear Descartes Hiring Team,</p>

  <p>I am applying for the <strong>.NET Software Developer (Data Integrations)</strong> role at Descartes Systems Group. With 15+ years of software engineering experience specializing in C#, backend data pipelines, high-volume transactional processing, and advanced SQL Server engineering, I have built a career around creating reliable data transformation systems and automated internal utilities for enterprise platforms.</p>

  <p>Descartes' mission to power global trade and logistics through reliable, high-volume data integrations aligns directly with my technical background and practical delivery experience:</p>

  <ul>
    <li><strong>Database &amp; SQL Server Engineering:</strong> Extensive experience designing complex stored procedures, triggers, user-defined functions, and ETL processes in Microsoft SQL Server, with a sharp focus on query optimization, index tuning, data modeling, and transaction integrity.</li>
    <li><strong>Full-Stack .NET &amp; Windows Worker Services:</strong> Hands-on history developing robust backend applications using modern .NET, Entity Framework, LINQ, ADO.NET, and self-healing Windows Worker Services integrated with Microsoft Azure cloud infrastructure.</li>
    <li><strong>Data Formats, Protocols &amp; B2B Integration:</strong> Extensive work handling high-throughput ingestion and transformation across XML, JSON, CSV, FTP/SFTP, REST APIs, and encrypted communication layers.</li>
    <li><strong>AI-First Engineering &amp; Quality Automation:</strong> Active advocate for AI-assisted development tools (Claude Code, GitHub Copilot) and automated testing (TDD, CI/CD with Azure DevOps and Jenkins) to compound delivery speed while safeguarding code quality.</li>
  </ul>

  <p>Based in Montevideo and fully prepared for your hybrid schedule (2 days per week in-office), I communicate with complete professional fluency in English and bring a collaborative, problem-solving mindset to the team.</p>

  <p>I look forward to discussing how my .NET and database integration experience can contribute to the success of Descartes' Data Integrations team.</p>

  <div class="signature">Sincerely,<br><strong>Pedro Milburn</strong></div>
</body>
</html>
"""

cl_html_path = app_dir / f"CL-{slug}.html"
cl_html_path.write_text(cl_html.strip(), encoding="utf-8")
print(f"Written CL HTML: {cl_html_path}")

# -------------------------------------------------------------
# 4. Cover Letter DOCX (English)
# -------------------------------------------------------------
cl_docx_path = app_dir / f"CL-{slug}.docx"
doc_cl = docx.Document()

for section in doc_cl.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style_cl = doc_cl.styles['Normal']
style_cl.font.name = 'Arial'
style_cl.font.size = Pt(10.5)
style_cl.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

p_cl_name = doc_cl.add_paragraph()
p_cl_name.paragraph_format.space_after = Pt(2)
r_cl_name = p_cl_name.add_run('Pedro Milburn')
r_cl_name.bold = True
r_cl_name.font.size = Pt(16)
r_cl_name.font.color.rgb = PRIMARY_COLOR

p_cl_contact = doc_cl.add_paragraph('Montevideo, Uruguay (Hybrid) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_cl_contact.runs[0].font.size = Pt(9.5)
p_cl_contact.runs[0].font.color.rgb = GRAY_COLOR
p_cl_contact.paragraph_format.space_after = Pt(14)

p_cl_date = doc_cl.add_paragraph('September 3, 2026')
p_cl_date.paragraph_format.space_after = Pt(12)

p_cl_rec = doc_cl.add_paragraph('Hiring Team / Data Integrations\nDescartes Systems Group (Montevideo)')
p_cl_rec.paragraph_format.space_after = Pt(14)

p_cl_sal = doc_cl.add_paragraph('Dear Descartes Hiring Team,')
p_cl_sal.paragraph_format.space_after = Pt(10)

p_cl_open = doc_cl.add_paragraph(
    'I am applying for the .NET Software Developer (Data Integrations) role at Descartes Systems Group. With 15+ years of software '
    'engineering experience specializing in C#, backend data pipelines, high-volume transactional processing, and advanced '
    'SQL Server engineering, I have built a career around creating reliable data transformation systems and automated internal '
    'utilities for enterprise platforms.'
)
p_cl_open.paragraph_format.space_after = Pt(10)
p_cl_open.paragraph_format.line_spacing = 1.15

p_cl_body = doc_cl.add_paragraph(
    'Descartes\' mission to power global trade and logistics through reliable, high-volume data integrations aligns directly with my technical background and practical delivery experience:'
)
p_cl_body.paragraph_format.space_after = Pt(8)
p_cl_body.paragraph_format.line_spacing = 1.15

cl_b1 = doc_cl.add_paragraph(style='List Bullet')
r1 = cl_b1.add_run('Database & SQL Server Engineering: ')
r1.bold = True
cl_b1.add_run(
    'Extensive experience designing complex stored procedures, triggers, user-defined functions, and ETL processes in Microsoft SQL Server, with a sharp focus on query optimization, index tuning, data modeling, and transaction integrity.'
)
cl_b1.paragraph_format.space_after = Pt(4)
cl_b1.paragraph_format.line_spacing = 1.15

cl_b2 = doc_cl.add_paragraph(style='List Bullet')
r2 = cl_b2.add_run('Full-Stack .NET & Windows Worker Services: ')
r2.bold = True
cl_b2.add_run(
    'Hands-on history developing robust backend applications using modern .NET, Entity Framework, LINQ, ADO.NET, and self-healing Windows Worker Services integrated with Microsoft Azure cloud infrastructure.'
)
cl_b2.paragraph_format.space_after = Pt(4)
cl_b2.paragraph_format.line_spacing = 1.15

cl_b3 = doc_cl.add_paragraph(style='List Bullet')
r3 = cl_b3.add_run('Data Formats, Protocols & B2B Integration: ')
r3.bold = True
cl_b3.add_run(
    'Extensive work handling high-throughput ingestion and transformation across XML, JSON, CSV, FTP/SFTP, REST APIs, and encrypted communication layers.'
)
cl_b3.paragraph_format.space_after = Pt(4)
cl_b3.paragraph_format.line_spacing = 1.15

cl_b4 = doc_cl.add_paragraph(style='List Bullet')
r4 = cl_b4.add_run('AI-First Engineering & Quality Automation: ')
r4.bold = True
cl_b4.add_run(
    'Active advocate for AI-assisted development tools (Claude Code, GitHub Copilot) and automated testing (TDD, CI/CD with Azure DevOps and Jenkins) to compound delivery speed while safeguarding code quality.'
)
cl_b4.paragraph_format.space_after = Pt(10)
cl_b4.paragraph_format.line_spacing = 1.15

p_cl_close1 = doc_cl.add_paragraph(
    'Based in Montevideo and fully prepared for your hybrid schedule (2 days per week in-office), I communicate with complete professional fluency in English and bring a collaborative, problem-solving mindset to the team.'
)
p_cl_close1.paragraph_format.space_after = Pt(10)
p_cl_close1.paragraph_format.line_spacing = 1.15

p_cl_close2 = doc_cl.add_paragraph('I look forward to discussing how my .NET and database integration experience can contribute to the success of Descartes\' Data Integrations team.')
p_cl_close2.paragraph_format.space_after = Pt(16)

p_cl_sign = doc_cl.add_paragraph('Sincerely,\n')
r_cl_sig = p_cl_sign.add_run('Pedro Milburn')
r_cl_sig.bold = True

doc_cl.save(str(cl_docx_path))
print(f"Written CL DOCX: {cl_docx_path}")

# -------------------------------------------------------------
# 5. Convert DOCX to PDF via Word COM
# -------------------------------------------------------------
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    # CV PDF
    cv_pdf_path = app_dir / f"CV-{slug}.pdf"
    doc1 = word.Documents.Open(str(cv_docx_path.resolve()))
    doc1.SaveAs(str(cv_pdf_path.resolve()), FileFormat=17)
    doc1.Close()
    print(f"Written CV PDF: {cv_pdf_path}")

    # CL PDF
    cl_pdf_path = app_dir / f"CL-{slug}.pdf"
    doc2 = word.Documents.Open(str(cl_docx_path.resolve()))
    doc2.SaveAs(str(cl_pdf_path.resolve()), FileFormat=17)
    doc2.Close()
    print(f"Written CL PDF: {cl_pdf_path}")
finally:
    word.Quit()

print("All Descartes application documents generated successfully!")
