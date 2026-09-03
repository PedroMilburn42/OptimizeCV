import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
import win32com.client

slug = "lawnstarter-software_engineering_manager-260903"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for Software Engineering Manager)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">Software Engineering Manager &amp; Technical Delivery Lead</div>
    <div class="contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>Software Engineering Manager with 15+ years of software engineering experience and a proven track record leading distributed, high-velocity delivery teams of 4 to 8 engineers. People-first leader focused on building individual career roadmaps, coaching senior engineers into tech leads, managing performance directly, and fostering autonomous "Managers of One." Pioneer in integrating AI-native workflows (spec-driven development, Claude Code, AI-assisted code review) to compound team throughput while protecting software craftsmanship.</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">People Leadership &amp; Growth:</span> 1:1 Mentorship, Individual IC Growth Roadmaps, Developing Tech Leads, Direct Conflict Resolution, Structured Performance Management, Technical Hiring &amp; Scorecards.</p>
  <p class="competency-row"><span class="competency-label">Delivery &amp; Operations:</span> Predictable Release Cadences, Async Decision-Making, Cycle Time &amp; Review Latency Optimization, Risk Surfacing, Cross-Functional Interface (Product, Design, QA, Architecture).</p>
  <p class="competency-row"><span class="competency-label">AI-Native Engineering:</span> Spec-Driven Development, Claude Code &amp; AI Agent Workflows, Automated Code Review Acceleration, MLOps Essentials (2026), AI Pair Programming (2026).</p>
  <p class="competency-row"><span class="competency-label">Technical Credibility:</span> Cloud Microservices, High-Throughput APIs, System Design Reviews, Distributed Architectures (GCP, Azure), Event-Driven Queues, SQL Server, PostgreSQL.</p>
  <p class="competency-row"><span class="competency-label">Engineering Hygiene &amp; Quality:</span> CI/CD Automation (Jenkins, Azure DevOps), TDD, SonarQube, Veracode, Incident Management &amp; Blameless RCAs.</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer &amp; Pod Lead | Aven Hospitality</span> <span>Jan 2026 - Present</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Remote/Hybrid)</div>
    <ul>
      <li>Lead a cross-functional pod of 8 engineers and QA specialists owning core transaction and booking systems end-to-end.</li>
      <li>Established individual career growth plans for every IC, conducting bi-weekly 1:1 coaching to develop emerging technical leads.</li>
      <li>Instituted lightweight async delivery cadences and AI-assisted code review workflows, reducing PR cycle times and eliminating delivery bottlenecks.</li>
      <li>Commanded critical incident responses with blameless RCAs, translating post-mortem findings into preventative automated testing guardrails.</li>
      <li>Partnered with Product and Operations to balance new marketplace feature velocity with architectural durability on GCP (Cloud Run, PubSub).</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Managed a team of 8 Senior engineers delivering core microservices and transactional business logic for a global travel marketplace.</li>
      <li>Created structured career progression frameworks, conducting regular performance evaluations and managing underperformance proactively.</li>
      <li>Monitored delivery health metrics (sprint velocity, review turnaround, error budgets), consistently meeting multi-quarter roadmap commitments.</li>
      <li>Served as primary interface between engineering, product managers, and enterprise architecture to make pragmatic technical trade-off decisions.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; Technical Lead | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led technical delivery and architecture for distributed multi-team pods supporting Fortune 500 accounts in the US and UK.</li>
      <li>Developed custom developer productivity tooling and early-generation AI automation plugins to accelerate team engineering workflows.</li>
      <li>Championed an async-first documentation culture in GitHub and Notion, enabling seamless collaboration across international time zones.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Team Lead | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led a team of 4 engineers refactoring enterprise platforms for UPM Uruguay, ensuring predictable release milestones and high test coverage.</li>
      <li>Mentored junior and mid-level developers on clean architecture, unit testing discipline, and automated build verification.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Led distributed engineering delivery for international clients (Volvo Cars Corporation, Envac AB, Collector Bank).</li>
      <li>Coached engineers on system reliability, asynchronous API integration, and clean code practices.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed high-throughput real-time bidding algorithms and multithreaded transaction backend pipelines in C# and ASP.NET.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; Leadership Credentials:</span> <span>MLOps Essentials (2026), Generative AI Pair Programming (2026), Microsoft Certified PL-400 &amp; PL-100</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('Software Engineering Manager & Technical Delivery Lead')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    'Software Engineering Manager with 15+ years of software engineering experience and a proven track record leading distributed, '
    'high-velocity delivery teams of 4 to 8 engineers. People-first leader focused on building individual career roadmaps, '
    'coaching senior engineers into tech leads, managing performance directly, and fostering autonomous "Managers of One." '
    'Pioneer in integrating AI-native workflows (spec-driven development, Claude Code, AI-assisted code review) to compound team '
    'throughput while protecting software craftsmanship.'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('People Leadership & Growth', '1:1 Mentorship, Individual IC Growth Roadmaps, Developing Tech Leads, Direct Conflict Resolution, Structured Performance Management, Technical Hiring & Scorecards.')
add_comp('Delivery & Operations', 'Predictable Release Cadences, Async Decision-Making, Cycle Time & Review Latency Optimization, Risk Surfacing, Cross-Functional Interface (Product, Design, QA, Architecture).')
add_comp('AI-Native Engineering', 'Spec-Driven Development, Claude Code & AI Agent Workflows, Automated Code Review Acceleration, MLOps Essentials (2026), AI Pair Programming (2026).')
add_comp('Technical Credibility', 'Cloud Microservices, High-Throughput APIs, System Design Reviews, Distributed Architectures (GCP, Azure), Event-Driven Queues, SQL Server, PostgreSQL.')
add_comp('Engineering Hygiene & Quality', 'CI/CD Automation (Jenkins, Azure DevOps), TDD, SonarQube, Veracode, Incident Management & Blameless RCAs.')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer & Pod Lead', 'Aven Hospitality', 'Jan 2026 - Present', 'Montevideo, Uruguay (Remote/Hybrid)',
    [
        'Lead a cross-functional pod of 8 engineers and QA specialists owning core transaction and booking systems end-to-end.',
        'Established individual career growth plans for every IC, conducting bi-weekly 1:1 coaching to develop emerging technical leads.',
        'Instituted lightweight async delivery cadences and AI-assisted code review workflows, reducing PR cycle times and eliminating delivery bottlenecks.',
        'Commanded critical incident responses with blameless RCAs, translating post-mortem findings into preventative automated testing guardrails.',
        'Partnered with Product and Operations to balance new marketplace feature velocity with architectural durability on GCP (Cloud Run, PubSub).'
    ]
)

add_role_cv(
    'Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Managed a team of 8 Senior engineers delivering core microservices and transactional business logic for a global travel marketplace.',
        'Created structured career progression frameworks, conducting regular performance evaluations and managing underperformance proactively.',
        'Monitored delivery health metrics (sprint velocity, review turnaround, error budgets), consistently meeting multi-quarter roadmap commitments.',
        'Served as primary interface between engineering, product managers, and enterprise architecture to make pragmatic technical trade-off decisions.'
    ]
)

add_role_cv(
    'Power Platform Architect & Technical Lead', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Led technical delivery and architecture for distributed multi-team pods supporting Fortune 500 accounts in the US and UK.',
        'Developed custom developer productivity tooling and early-generation AI automation plugins to accelerate team engineering workflows.',
        'Championed an async-first documentation culture in GitHub and Notion, enabling seamless collaboration across international time zones.'
    ]
)

add_role_cv(
    'Software Architect & Team Lead', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led a team of 4 engineers refactoring enterprise platforms for UPM Uruguay, ensuring predictable release milestones and high test coverage.',
        'Mentored junior and mid-level developers on clean architecture, unit testing discipline, and automated build verification.'
    ]
)

add_role_cv(
    'Technical Solution Architect', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Led distributed engineering delivery for international clients (Volvo Cars Corporation, Envac AB, Collector Bank).',
        'Coached engineers on system reliability, asynchronous API integration, and clean code practices.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Developed high-throughput real-time bidding algorithms and multithreaded transaction backend pipelines in C# and ASP.NET.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('AI & Leadership Credentials: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('MLOps Essentials (2026), Generative AI Pair Programming (2026), Microsoft Certified PL-400 & PL-100')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('International Patent: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")

# -------------------------------------------------------------
# 3. Cover Letter HTML (English, addressed to LawnStarter)
# -------------------------------------------------------------
cl_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - Cover Letter</title>
<style>
  @page { size: A4; margin: 2.2cm 2.4cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 11pt; color: #1a1a1a; background: #ffffff; line-height: 1.45; }

  .sender { border-bottom: 2pt solid #1f3a5f; padding-bottom: 6pt; margin-bottom: 14pt; }
  .sender-name { font-size: 15pt; font-weight: bold; color: #1f3a5f; }
  .sender-contact { font-size: 9.5pt; color: #444; margin-top: 2pt; }

  .date { margin-bottom: 14pt; font-size: 10pt; color: #555; }
  .recipient { margin-bottom: 14pt; font-size: 10.5pt; }
  p { margin: 0 0 9pt 0; }
  ul { margin: 0 0 9pt 0; padding-left: 16pt; break-inside: avoid; }
  li { margin-bottom: 3pt; }
  .signature { margin-top: 14pt; }
</style>
</head>
<body>
  <div class="sender">
    <div class="sender-name">Pedro Milburn</div>
    <div class="sender-contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>
  <div class="date">September 3, 2026</div>
  <div class="recipient">Head of Engineering &amp; Hiring Team<br>LawnStarter</div>

  <p>Dear LawnStarter Hiring Team,</p>

  <p>I am applying for the <strong>Software Engineering Manager</strong> role at LawnStarter. Having led high-velocity, distributed engineering pods of 4 to 8 engineers across enterprise marketplaces and SaaS platforms, I have built my management philosophy around people-first leadership: coaching engineers into high-ownership "Managers of One," running lightweight cadences, and actively accelerating delivery through AI-native workflows.</p>

  <p>LawnStarter's focus on scaling new service verticals while maintaining delivery predictability resonates strongly with my experience leading engineering pods at Aven Hospitality and Sabre Corporation:</p>

  <ul>
    <li><strong>People Development &amp; Direct Leadership:</strong> I build documented growth plans for every engineer on my team, using 1:1s for career progression rather than status theater. I address friction and underperformance directly through clear, empathetic processes while mentoring senior engineers toward technical lead responsibilities.</li>
    <li><strong>Predictable Delivery &amp; Cross-Functional Interface:</strong> I run lightweight agile cadences that keep Product, Design, and Engineering aligned. By tracking delivery metrics (cycle time, PR review turnaround, sprint predictability) and spotting risks early, my pods deliver commitments on cadence without becoming bottlenecks.</li>
    <li><strong>AI-Native Engineering Acceleration:</strong> I actively champion spec-driven development and AI agent workflows (Claude Code, automated PR review, ticket sharpening) to eliminate manual overhead, holding MLOps and Generative AI certifications (2026) to compound team throughput while protecting software quality.</li>
  </ul>

  <p>Operating as a remote-native leader based in Montevideo, Uruguay, I communicate with clarity across async channels and collaborate seamlessly across time zones to drive outcomes.</p>

  <p>I look forward to discussing how I can help LawnStarter scale its engineering capacity and grow its next generation of technical leaders.</p>

  <div class="signature">Sincerely,<br><strong>Pedro Milburn</strong></div>
</body>
</html>
"""

cl_html_path = app_dir / f"CL-{slug}.html"
cl_html_path.write_text(cl_html.strip(), encoding="utf-8")
print(f"Written CL HTML: {cl_html_path}")

# -------------------------------------------------------------
# 4. Cover Letter DOCX (English)
# -------------------------------------------------------------
cl_docx_path = app_dir / f"CL-{slug}.docx"
doc_cl = docx.Document()

for section in doc_cl.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

style_cl = doc_cl.styles['Normal']
style_cl.font.name = 'Arial'
style_cl.font.size = Pt(10.5)
style_cl.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

p_cl_name = doc_cl.add_paragraph()
p_cl_name.paragraph_format.space_after = Pt(2)
r_cl_name = p_cl_name.add_run('Pedro Milburn')
r_cl_name.bold = True
r_cl_name.font.size = Pt(16)
r_cl_name.font.color.rgb = PRIMARY_COLOR

p_cl_contact = doc_cl.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_cl_contact.runs[0].font.size = Pt(9.5)
p_cl_contact.runs[0].font.color.rgb = GRAY_COLOR
p_cl_contact.paragraph_format.space_after = Pt(14)

p_cl_date = doc_cl.add_paragraph('September 3, 2026')
p_cl_date.paragraph_format.space_after = Pt(12)

p_cl_rec = doc_cl.add_paragraph('Head of Engineering & Hiring Team\nLawnStarter')
p_cl_rec.paragraph_format.space_after = Pt(14)

p_cl_sal = doc_cl.add_paragraph('Dear LawnStarter Hiring Team,')
p_cl_sal.paragraph_format.space_after = Pt(10)

p_cl_open = doc_cl.add_paragraph(
    'I am applying for the Software Engineering Manager role at LawnStarter. Having led high-velocity, distributed '
    'engineering pods of 4 to 8 engineers across enterprise marketplaces and SaaS platforms, I have built my management '
    'philosophy around people-first leadership: coaching engineers into high-ownership "Managers of One," running '
    'lightweight cadences, and actively accelerating delivery through AI-native workflows.'
)
p_cl_open.paragraph_format.space_after = Pt(10)
p_cl_open.paragraph_format.line_spacing = 1.15

p_cl_body = doc_cl.add_paragraph(
    'LawnStarter\'s focus on scaling new service verticals while maintaining delivery predictability resonates strongly with my experience leading engineering pods at Aven Hospitality and Sabre Corporation:'
)
p_cl_body.paragraph_format.space_after = Pt(8)
p_cl_body.paragraph_format.line_spacing = 1.15

cl_b1 = doc_cl.add_paragraph(style='List Bullet')
r1 = cl_b1.add_run('People Development & Direct Leadership: ')
r1.bold = True
cl_b1.add_run(
    'I build documented growth plans for every engineer on my team, using 1:1s for career progression rather than status theater. '
    'I address friction and underperformance directly through clear, empathetic processes while mentoring senior engineers toward technical lead responsibilities.'
)
cl_b1.paragraph_format.space_after = Pt(4)
cl_b1.paragraph_format.line_spacing = 1.15

cl_b2 = doc_cl.add_paragraph(style='List Bullet')
r2 = cl_b2.add_run('Predictable Delivery & Cross-Functional Interface: ')
r2.bold = True
cl_b2.add_run(
    'I run lightweight agile cadences that keep Product, Design, and Engineering aligned. By tracking delivery metrics (cycle time, '
    'PR review turnaround, sprint predictability) and spotting risks early, my pods deliver commitments on cadence without becoming bottlenecks.'
)
cl_b2.paragraph_format.space_after = Pt(4)
cl_b2.paragraph_format.line_spacing = 1.15

cl_b3 = doc_cl.add_paragraph(style='List Bullet')
r3 = cl_b3.add_run('AI-Native Engineering Acceleration: ')
r3.bold = True
cl_b3.add_run(
    'I actively champion spec-driven development and AI agent workflows (Claude Code, automated PR review, ticket sharpening) to eliminate '
    'manual overhead, holding MLOps and Generative AI certifications (2026) to compound team throughput while protecting software quality.'
)
cl_b3.paragraph_format.space_after = Pt(10)
cl_b3.paragraph_format.line_spacing = 1.15

p_cl_close1 = doc_cl.add_paragraph(
    'Operating as a remote-native leader based in Montevideo, Uruguay, I communicate with clarity across async channels '
    'and collaborate seamlessly across time zones to drive outcomes.'
)
p_cl_close1.paragraph_format.space_after = Pt(10)
p_cl_close1.paragraph_format.line_spacing = 1.15

p_cl_close2 = doc_cl.add_paragraph('I look forward to discussing how I can help LawnStarter scale its engineering capacity and grow its next generation of technical leaders.')
p_cl_close2.paragraph_format.space_after = Pt(16)

p_cl_sign = doc_cl.add_paragraph('Sincerely,\n')
r_cl_sig = p_cl_sign.add_run('Pedro Milburn')
r_cl_sig.bold = True

doc_cl.save(str(cl_docx_path))
print(f"Written CL DOCX: {cl_docx_path}")

# -------------------------------------------------------------
# 5. Convert DOCX to PDF via Word COM
# -------------------------------------------------------------
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    # CV PDF
    cv_pdf_path = app_dir / f"CV-{slug}.pdf"
    doc1 = word.Documents.Open(str(cv_docx_path.resolve()))
    doc1.SaveAs(str(cv_pdf_path.resolve()), FileFormat=17)
    doc1.Close()
    print(f"Written CV PDF: {cv_pdf_path}")

    # CL PDF
    cl_pdf_path = app_dir / f"CL-{slug}.pdf"
    doc2 = word.Documents.Open(str(cl_docx_path.resolve()))
    doc2.SaveAs(str(cl_pdf_path.resolve()), FileFormat=17)
    doc2.Close()
    print(f"Written CL PDF: {cl_pdf_path}")
finally:
    word.Quit()

print("All LawnStarter application documents generated successfully!")
