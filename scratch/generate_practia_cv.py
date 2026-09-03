import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor

slug = "practia-senior_dotnet_software_engineer-260831"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for Practia .NET 8 & Architecture)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">Senior .NET Software Engineer &amp; Solutions Architect | Former Practia Consultant</div>
    <div class="contact">Montevideo, Uruguay &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>Senior .NET Software Engineer and Solutions Architect with 15+ years designing and implementing high-impact enterprise applications, cloud microservices, and distributed backend systems. Former Practia consultant with a proven track record architecting solutions for major corporate and financial clients. Expert in .NET 8, Clean Architecture, Domain-Driven Design (DDD / ABP.IO framework paradigms), CQRS, Event-Driven Messaging, and API Gateways.</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">.NET &amp; Architecture:</span> C#, .NET 8, .NET Core, ASP.NET Core, Clean Architecture, Domain-Driven Design (DDD), ABP.IO concepts, CQRS, Microservices, Saga / Distributed Transactions.</p>
  <p class="competency-row"><span class="competency-label">APIs &amp; Integration:</span> RESTful APIs, API Gateway, API Management, GraphQL, Message Queues (PubSub, Event-Driven), WCF, SOAP.</p>
  <p class="competency-row"><span class="competency-label">Databases &amp; Cache:</span> SQL Server (T-SQL, optimization, stored procedures), PostgreSQL, Oracle, Redis / Caching, Entity Framework Core, Dapper.</p>
  <p class="competency-row"><span class="competency-label">Cloud &amp; DevOps:</span> Microsoft Azure (AKS, Functions, App Services), GCP (Cloud Run, VMs), Docker, CI/CD (Azure DevOps, Jenkins, TeamCity), Git.</p>
  <p class="competency-row"><span class="competency-label">Frontend &amp; Engineering Quality:</span> Angular, JavaScript/TypeScript, TDD, Unit &amp; Integration Testing, SonarQube, Veracode, Agile / Scrum.</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer | Aven Hospitality</span> <span>Jan 2026 - Present</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Hybrid)</div>
    <ul>
      <li>Lead a multidisciplinary team of 8 engineers and QA specialists developing high-throughput booking backend services in modern .NET.</li>
      <li>Architected decoupled event-driven messaging pipelines and queue processing on GCP (Cloud Run, PubSub) for distributed transactions.</li>
      <li>Automated CI/CD deployment pipelines via Jenkins, Docker, and PowerShell scripts, maintaining zero-downtime releases.</li>
      <li>Enforce code reviews, SOLID principles, and security remediation using Veracode and SonarQube.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led an engineering pod of 8 Senior .NET developers delivering core transaction logic and microservices on Microsoft Azure.</li>
      <li>Architected distributed cloud services (AKS, Azure Functions, App Services), resolving critical bottlenecks in transactional latency.</li>
      <li>Standardized Clean Architecture guidelines, peer code reviews, and Azure DevOps release automation to improve delivery velocity.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; .NET Developer | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Architected and delivered enterprise cloud solutions, custom .NET extensions, and API integrations for Fortune 500 accounts in the US and UK.</li>
      <li>Engineered developer tooling in C# and .NET integrated with cloud platforms and internal automation systems.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Consultant | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led a team of 4 developers refactoring core enterprise systems for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server.</li>
      <li>Designed and implemented enterprise portal solutions and Business Intelligence platforms for major regulated clients, including Banco de Chile.</li>
      <li>Implemented continuous integration and automated build testing on TeamCity build servers.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>.NET Developer | Urudata Software (Banco Itaú Project)</span> <span>Sep 2016 - Dec 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Developed central document digitalization and teller tracking software in .NET MVC 5 and SQL Server across 20 Banco Itaú branches.</li>
      <li>Integrated business process automation engines (Q-Flow) under strict banking security and regulatory standards.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect &amp; .NET Developer | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed Azure backend REST services and payment gateway integrations with AngularJS frontend for Collector Bank mobile payments.</li>
      <li>Engineered high-availability multithreaded Windows Services and cross-platform APIs with SQL stored procedures for Envac AB global operations.</li>
      <li>Led core architectural delivery for Volvo Cars Corporation enterprise intranet systems on Microsoft platforms.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Developed multithreaded Windows Services, optimization algorithms, and high-volume transaction processing systems in C# and ASP.NET.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>Microsoft Certifications:</span> <span>Power Platform Developer Associate (PL-400), App Maker (PL-100), MCTS C# &amp; SharePoint</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; MLOps:</span> <span>MLOps Essentials (2026), Generative AI Pair Programming (2026), Azure DevOps</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('Senior .NET Software Engineer & Solutions Architect | Former Practia Consultant')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    'Senior .NET Software Engineer and Solutions Architect with 15+ years designing and implementing high-impact enterprise applications, '
    'cloud microservices, and distributed backend systems. Former Practia consultant with a proven track record architecting solutions '
    'for major corporate and financial clients. Expert in .NET 8, Clean Architecture, Domain-Driven Design (DDD / ABP.IO framework paradigms), '
    'CQRS, Event-Driven Messaging, and API Gateways.'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('.NET & Architecture', 'C#, .NET 8, .NET Core, ASP.NET Core, Clean Architecture, Domain-Driven Design (DDD), ABP.IO concepts, CQRS, Microservices, Saga / Distributed Transactions.')
add_comp('APIs & Integration', 'RESTful APIs, API Gateway, API Management, GraphQL, Message Queues (PubSub, Event-Driven), WCF, SOAP.')
add_comp('Databases & Cache', 'SQL Server (T-SQL, optimization, stored procedures), PostgreSQL, Oracle, Redis / Caching, Entity Framework Core, Dapper.')
add_comp('Cloud & DevOps', 'Microsoft Azure (AKS, Functions, App Services), GCP (Cloud Run, VMs), Docker, CI/CD (Azure DevOps, Jenkins, TeamCity), Git.')
add_comp('Frontend & Engineering Quality', 'Angular, JavaScript/TypeScript, TDD, Unit & Integration Testing, SonarQube, Veracode, Agile / Scrum.')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer', 'Aven Hospitality', 'Jan 2026 - Present', 'Montevideo, Uruguay (Hybrid)',
    [
        'Lead a multidisciplinary team of 8 engineers and QA specialists developing high-throughput booking backend services in modern .NET.',
        'Architected decoupled event-driven messaging pipelines and queue processing on GCP (Cloud Run, PubSub) for distributed transactions.',
        'Automated CI/CD deployment pipelines via Jenkins, Docker, and PowerShell scripts, maintaining zero-downtime releases.',
        'Enforce code reviews, SOLID principles, and security remediation using Veracode and SonarQube.'
    ]
)

add_role_cv(
    'Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Led an engineering pod of 8 Senior .NET developers delivering core transaction logic and microservices on Microsoft Azure.',
        'Architected distributed cloud services (AKS, Azure Functions, App Services), resolving critical bottlenecks in transactional latency.',
        'Standardized Clean Architecture guidelines, peer code reviews, and Azure DevOps release automation to improve delivery velocity.'
    ]
)

add_role_cv(
    'Power Platform Architect & .NET Developer', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Architected and delivered enterprise cloud solutions, custom .NET extensions, and API integrations for Fortune 500 accounts in the US and UK.',
        'Engineered developer tooling in C# and .NET integrated with cloud platforms and internal automation systems.'
    ]
)

add_role_cv(
    'Software Architect & Consultant', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led a team of 4 developers refactoring core enterprise systems for UPM Uruguay into a decoupled 3-layer architecture with Entity Framework and SQL Server.',
        'Designed and implemented enterprise portal solutions and Business Intelligence platforms for major regulated clients, including Banco de Chile.',
        'Implemented continuous integration and automated build testing on TeamCity build servers.'
    ]
)

add_role_cv(
    '.NET Developer', 'Urudata Software (Banco Itaú Project)', 'Sep 2016 - Dec 2016', 'Montevideo, Uruguay',
    [
        'Developed central document digitalization and teller tracking software in .NET MVC 5 and SQL Server across 20 Banco Itaú branches.',
        'Integrated business process automation engines (Q-Flow) under strict banking security and regulatory standards.'
    ]
)

add_role_cv(
    'Technical Solution Architect & .NET Developer', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Developed Azure backend REST services and payment gateway integrations with AngularJS frontend for Collector Bank mobile payments.',
        'Engineered high-availability multithreaded Windows Services and cross-platform APIs with SQL stored procedures for Envac AB global operations.',
        'Led core architectural delivery for Volvo Cars Corporation enterprise intranet systems on Microsoft platforms.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Developed multithreaded Windows Services, optimization algorithms, and high-volume transaction processing systems in C# and ASP.NET.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('Microsoft Certifications: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('Power Platform Developer Associate (PL-400), App Maker (PL-100), MCTS C# & SharePoint')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('AI & MLOps: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('MLOps Essentials (2026), Generative AI Pair Programming (2026), Azure DevOps')

pc3 = doc_cv.add_paragraph()
pc3.paragraph_format.space_after = Pt(2)
rc3_l = pc3.add_run('International Patent: ')
rc3_l.bold = True
rc3_l.font.color.rgb = PRIMARY_COLOR
pc3.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")
