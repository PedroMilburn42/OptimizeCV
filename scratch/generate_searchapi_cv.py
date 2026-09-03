import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor

slug = "searchapi-engineering_manager-260831"
app_dir = Path("data/applications") / slug
app_dir.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 1. CV HTML (English, Tailored for Engineering Manager)
# -------------------------------------------------------------
cv_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>Pedro Milburn - CV</title>
<style>
  @page { size: A4; margin: 1.4cm 1.7cm; }
  body { font-family: "Helvetica Neue", Arial, sans-serif; font-size: 9.8pt; color: #1a1a1a; background: #ffffff; line-height: 1.32; }

  .header { border-bottom: 2pt solid #1f3a5f; padding-bottom: 5pt; margin-bottom: 7pt; }
  h1 { font-size: 19pt; color: #1f3a5f; margin: 0 0 1pt 0; }
  .tagline { font-size: 10.5pt; color: #444; font-style: italic; margin-bottom: 3pt; }
  .contact { font-size: 9pt; color: #444; }

  h2 { font-size: 10.5pt; color: #1f3a5f; text-transform: uppercase; letter-spacing: 0.5pt; border-bottom: 1pt solid #1f3a5f; margin: 9pt 0 4pt 0; break-after: avoid; }

  .competency-row { margin-bottom: 2.5pt; break-inside: avoid; font-size: 9.2pt; }
  .competency-label { font-weight: bold; color: #1f3a5f; }

  .entry { margin-bottom: 6.5pt; break-inside: avoid; }
  .entry-header { display: flex; justify-content: space-between; font-weight: bold; font-size: 9.8pt; }
  .entry-meta { font-size: 8.8pt; color: #555; font-style: italic; margin-bottom: 1.5pt; }
  ul { margin: 1.5pt 0 0 0; padding-left: 13pt; }
  li { margin-bottom: 1.5pt; font-size: 9.2pt; }
</style>
</head>
<body>
  <div class="header">
    <h1>Pedro Milburn</h1>
    <div class="tagline">Engineering Manager &amp; Technical Delivery Lead</div>
    <div class="contact">Montevideo, Uruguay (Remote) &middot; pierre.milburn@hotmail.com &middot; +598 92 806 401 &middot; linkedin.com/in/pedro-milburn-a0700417</div>
  </div>

  <h2>Profile</h2>
  <p>Engineering Manager and Technical Delivery Lead with 15+ years guiding engineering teams, establishing high-trust operational cadences, and taking full ownership of product delivery. Experienced in managing cross-functional pods of 8+ engineers and QA, commanding production incidents with blameless post-mortems (RCAs), unblocking delivery bottlenecks, and leveraging AI agent workflows to accelerate engineering processes without compromising on quality or architectural rigor.</p>

  <h2>Core Competencies</h2>
  <p class="competency-row"><span class="competency-label">Engineering Management:</span> End-to-End Delivery Ownership, Agile Cadences (Planning, Standups, Retros), 1:1 Coaching, Hiring &amp; Technical Screening, Building "Managers of One".</p>
  <p class="competency-row"><span class="competency-label">Delivery Metrics &amp; Operations:</span> Cycle Time, Review Latency, SLA Management, Incident Command, Root Cause Analysis (Blameless RCAs), Risk Mitigation.</p>
  <p class="competency-row"><span class="competency-label">Technical Mastery:</span> High-Throughput RESTful APIs, Distributed Systems, Event-Driven Queues (PubSub, Message Queues), Cloud Architecture (GCP, Azure), SQL Server.</p>
  <p class="competency-row"><span class="competency-label">DevOps &amp; Automation:</span> GitHub, Notion, CI/CD Automation (Jenkins, Azure DevOps), Docker, Automated Testing (TDD), SonarQube, Veracode.</p>
  <p class="competency-row"><span class="competency-label">AI &amp; Workflow Acceleration:</span> AI Agent Integration, Process Automation, MLOps Essentials (2026), Generative AI Pair Programming (2026).</p>

  <h2>Professional Experience</h2>
  <div class="entry">
    <div class="entry-header"><span>Lead Software Engineer &amp; Pod Lead | Aven Hospitality</span> <span>Jan 2026 - Present</span></div>
    <div class="entry-meta">Montevideo, Uruguay (Remote/Hybrid)</div>
    <ul>
      <li>Own end-to-end delivery for core transaction and booking platforms, managing an engineering pod of 8 developers and QA specialists.</li>
      <li>Designed and ran lightweight agile cadences, streamlining sprint planning and reducing PR review turnaround times.</li>
      <li>Commanded production incident responses, authoring blameless RCAs and implementing preventative automated test guardrails.</li>
      <li>Modernized backend architectures to GCP (Cloud Run, PubSub) and automated deployment pipelines via Jenkins and Docker.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Development Team Lead | Sabre Corporation</span> <span>Oct 2023 - Feb 2026</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Managed an engineering team of 8 Senior .NET engineers delivering core microservices and business logic for a global travel marketplace.</li>
      <li>Monitored delivery velocity, cycle times, and operational health, removing cross-team blockers to maintain on-time release milestones.</li>
      <li>Served as first escalation point for critical production outages, coordinating cross-functional triage and resolving latency bottlenecks on Azure.</li>
      <li>Conducted 1:1 coaching, technical performance reviews, and spearheaded hiring scorecards to maintain a high talent bar.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Power Platform Architect &amp; Technical Lead | Globant</span> <span>Sep 2018 - Sep 2023</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led technical delivery and architecture for enterprise client engagements across the US and UK, managing multi-team delivery pods.</li>
      <li>Built internal developer productivity tooling and integrated AI-assisted coding assistants into daily engineering workflows.</li>
      <li>Standardized asynchronous documentation in GitHub and Notion, ensuring seamless handoffs and alignment across distributed pods.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Architect &amp; Team Lead | Grupo Pragma Consultores (Practia)</span> <span>Aug 2015 - Aug 2016</span></div>
    <div class="entry-meta">Montevideo, Uruguay</div>
    <ul>
      <li>Led a team of 4 developers refactoring core enterprise platforms for UPM Uruguay, ensuring strict delivery milestones and test coverage.</li>
      <li>Implemented continuous integration and automated quality gates on TeamCity build servers.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Technical Solution Architect | HiQ</span> <span>Jan 2011 - Jul 2015</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Led engineering delivery for global clients (Volvo Cars Corporation, Envac AB, Collector Bank), designing resilient distributed services.</li>
      <li>Mentored developers on clean code standards, automated testing, and asynchronous API integrations.</li>
    </ul>
  </div>

  <div class="entry">
    <div class="entry-header"><span>Software Engineer | Admeta</span> <span>Feb 2007 - Dec 2009</span></div>
    <div class="entry-meta">Gothenburg, Sweden</div>
    <ul>
      <li>Engineered high-throughput real-time bidding algorithms and multithreaded backend data pipelines for online ad exchanges.</li>
    </ul>
  </div>

  <h2>Education &amp; Credentials</h2>
  <div class="entry">
    <div class="entry-header"><span>Master's Degree in Mobile Services | Chalmers University of Technology (Sweden)</span> <span>2004 - 2006</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>Degree Studies, Systems Engineering | Universidad ORT Uruguay</span> <span>1993 - 1999</span></div>
  </div>

  <h2>Certifications &amp; Patents</h2>
  <div class="entry">
    <div class="entry-header"><span>AI &amp; Engineering Credentials:</span> <span>MLOps Essentials (2026), Generative AI Pair Programming (2026), Microsoft Certified PL-400 &amp; PL-100</span></div>
  </div>
  <div class="entry">
    <div class="entry-header"><span>International Patent:</span> <span>Method for optimising the placement of at least two pieces of information (US 13/254,248)</span></div>
  </div>
</body>
</html>
"""

cv_html_path = app_dir / f"CV-{slug}.html"
cv_html_path.write_text(cv_html.strip(), encoding="utf-8")
print(f"Written CV HTML: {cv_html_path}")

# -------------------------------------------------------------
# 2. CV DOCX
# -------------------------------------------------------------
cv_docx_path = app_dir / f"CV-{slug}.docx"
doc_cv = docx.Document()

for section in doc_cv.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

style_cv = doc_cv.styles['Normal']
style_cv.font.name = 'Arial'
style_cv.font.size = Pt(9.5)
style_cv.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

PRIMARY_COLOR = RGBColor(0x1f, 0x3a, 0x5f)
GRAY_COLOR = RGBColor(0x55, 0x55, 0x55)

def add_cv_h2(text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = PRIMARY_COLOR
    return p

# Header
p_name = doc_cv.add_paragraph()
p_name.paragraph_format.space_after = Pt(1)
r_name = p_name.add_run('Pedro Milburn')
r_name.bold = True
r_name.font.size = Pt(18)
r_name.font.color.rgb = PRIMARY_COLOR

p_tag = doc_cv.add_paragraph()
p_tag.paragraph_format.space_after = Pt(2)
r_tag = p_tag.add_run('Engineering Manager & Technical Delivery Lead')
r_tag.italic = True
r_tag.font.size = Pt(10.5)
r_tag.font.color.rgb = GRAY_COLOR

p_contact = doc_cv.add_paragraph('Montevideo, Uruguay (Remote) · pierre.milburn@hotmail.com · +598 92 806 401 · linkedin.com/in/pedro-milburn-a0700417')
p_contact.runs[0].font.size = Pt(9)
p_contact.runs[0].font.color.rgb = GRAY_COLOR
p_contact.paragraph_format.space_after = Pt(6)

# Profile
add_cv_h2('Profile')
p_prof = doc_cv.add_paragraph(
    'Engineering Manager and Technical Delivery Lead with 15+ years guiding engineering teams, establishing high-trust '
    'operational cadences, and taking full ownership of product delivery. Experienced in managing cross-functional pods '
    'of 8+ engineers and QA, commanding production incidents with blameless post-mortems (RCAs), unblocking delivery bottlenecks, '
    'and leveraging AI agent workflows to accelerate engineering processes without compromising on quality or architectural rigor.'
)
p_prof.paragraph_format.space_after = Pt(4)
p_prof.paragraph_format.line_spacing = 1.15

# Competencies
add_cv_h2('Core Competencies')

def add_comp(label, text):
    p = doc_cv.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_l = p.add_run(f'{label}: ')
    r_l.bold = True
    r_l.font.color.rgb = PRIMARY_COLOR
    p.add_run(text)

add_comp('Engineering Management', 'End-to-End Delivery Ownership, Agile Cadences (Planning, Standups, Retros), 1:1 Coaching, Hiring & Technical Screening, Building "Managers of One".')
add_comp('Delivery Metrics & Operations', 'Cycle Time, Review Latency, SLA Management, Incident Command, Root Cause Analysis (Blameless RCAs), Risk Mitigation.')
add_comp('Technical Mastery', 'High-Throughput RESTful APIs, Distributed Systems, Event-Driven Queues (PubSub, Message Queues), Cloud Architecture (GCP, Azure), SQL Server.')
add_comp('DevOps & Automation', 'GitHub, Notion, CI/CD Automation (Jenkins, Azure DevOps), Docker, Automated Testing (TDD), SonarQube, Veracode.')
add_comp('AI & Workflow Acceleration', 'AI Agent Integration, Process Automation, MLOps Essentials (2026), Generative AI Pair Programming (2026).')

# Experience
add_cv_h2('Professional Experience')

def add_role_cv(title, company, dates, loc, bullets):
    p_head = doc_cv.add_paragraph()
    p_head.paragraph_format.space_before = Pt(4)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    r_t = p_head.add_run(f'{title} | {company}')
    r_t.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR
    r_d = p_head.add_run(f'   ({dates})')
    r_d.font.color.rgb = GRAY_COLOR
    r_d.italic = True
    
    p_loc = doc_cv.add_paragraph(loc)
    p_loc.runs[0].font.size = Pt(8.5)
    p_loc.runs[0].font.italic = True
    p_loc.runs[0].font.color.rgb = GRAY_COLOR
    p_loc.paragraph_format.space_after = Pt(2)
    p_loc.paragraph_format.keep_with_next = True
    
    for b in bullets:
        bp = doc_cv.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(1.5)
        bp.paragraph_format.line_spacing = 1.15
        bp.add_run(b)

add_role_cv(
    'Lead Software Engineer & Pod Lead', 'Aven Hospitality', 'Jan 2026 - Present', 'Montevideo, Uruguay (Remote/Hybrid)',
    [
        'Own end-to-end delivery for core transaction and booking platforms, managing an engineering pod of 8 developers and QA specialists.',
        'Designed and ran lightweight agile cadences, streamlining sprint planning and reducing PR review turnaround times.',
        'Commanded production incident responses, authoring blameless RCAs and implementing preventative automated test guardrails.',
        'Modernized backend architectures to GCP (Cloud Run, PubSub) and automated deployment pipelines via Jenkins and Docker.'
    ]
)

add_role_cv(
    'Development Team Lead', 'Sabre Corporation', 'Oct 2023 - Feb 2026', 'Montevideo, Uruguay',
    [
        'Managed an engineering team of 8 Senior .NET engineers delivering core microservices and business logic for a global travel marketplace.',
        'Monitored delivery velocity, cycle times, and operational health, removing cross-team blockers to maintain on-time release milestones.',
        'Served as first escalation point for critical production outages, coordinating cross-functional triage and resolving latency bottlenecks on Azure.',
        'Conducted 1:1 coaching, technical performance reviews, and spearheaded hiring scorecards to maintain a high talent bar.'
    ]
)

add_role_cv(
    'Power Platform Architect & Technical Lead', 'Globant', 'Sep 2018 - Sep 2023', 'Montevideo, Uruguay',
    [
        'Led technical delivery and architecture for enterprise client engagements across the US and UK, managing multi-team delivery pods.',
        'Built internal developer productivity tooling and integrated AI-assisted coding assistants into daily engineering workflows.',
        'Standardized asynchronous documentation in GitHub and Notion, ensuring seamless handoffs and alignment across distributed pods.'
    ]
)

add_role_cv(
    'Software Architect & Team Lead', 'Grupo Pragma Consultores (Practia)', 'Aug 2015 - Aug 2016', 'Montevideo, Uruguay',
    [
        'Led a team of 4 developers refactoring core enterprise platforms for UPM Uruguay, ensuring strict delivery milestones and test coverage.',
        'Implemented continuous integration and automated quality gates on TeamCity build servers.'
    ]
)

add_role_cv(
    'Technical Solution Architect', 'HiQ', 'Jan 2011 - Jul 2015', 'Gothenburg, Sweden',
    [
        'Led engineering delivery for global clients (Volvo Cars Corporation, Envac AB, Collector Bank), designing resilient distributed services.',
        'Mentored developers on clean code standards, automated testing, and asynchronous API integrations.'
    ]
)

add_role_cv(
    'Software Engineer', 'Admeta', 'Feb 2007 - Dec 2009', 'Gothenburg, Sweden',
    [
        'Engineered high-throughput real-time bidding algorithms and multithreaded backend data pipelines for online ad exchanges.'
    ]
)

# Education
add_cv_h2('Education & Credentials')
pe1 = doc_cv.add_paragraph()
pe1.paragraph_format.space_after = Pt(2)
re1_t = pe1.add_run('Master\'s Degree in Mobile Services | Chalmers University of Technology (Sweden)')
re1_t.bold = True
pe1.add_run(' (2004 - 2006)')

pe2 = doc_cv.add_paragraph()
pe2.paragraph_format.space_after = Pt(4)
re2_t = pe2.add_run('Degree Studies, Systems Engineering | Universidad ORT Uruguay')
re2_t.bold = True
pe2.add_run(' (1993 - 1999)')

# Certs
add_cv_h2('Certifications & Patents')
pc1 = doc_cv.add_paragraph()
pc1.paragraph_format.space_after = Pt(2)
rc1_l = pc1.add_run('AI & Engineering Credentials: ')
rc1_l.bold = True
rc1_l.font.color.rgb = PRIMARY_COLOR
pc1.add_run('MLOps Essentials (2026), Generative AI Pair Programming (2026), Microsoft Certified PL-400 & PL-100')

pc2 = doc_cv.add_paragraph()
pc2.paragraph_format.space_after = Pt(2)
rc2_l = pc2.add_run('International Patent: ')
rc2_l.bold = True
rc2_l.font.color.rgb = PRIMARY_COLOR
pc2.add_run('Method for optimising the placement of at least two pieces of information (US 13/254,248)')

doc_cv.save(str(cv_docx_path))
print(f"Written CV DOCX: {cv_docx_path}")
